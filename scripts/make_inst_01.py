import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from scripts.generate_all_instruments import (
    create_base_doc, add_header_kop, add_doc_title, add_section_heading,
    add_body_p, format_cell, set_table_borders, OUT_DIR
)

def build_instrument_01():
    doc = create_base_doc()
    add_header_kop(doc)
    add_doc_title(
        doc,
        "PEDOMAN WAWANCARA ANALISIS KEBUTUHAN GURU MATEMATIKA",
        "Penelitian Pengembangan Web Application Berbasis Generative Artificial Intelligence untuk Menghasilkan LKPD Matematika Berdiferensiasi Berdasarkan Kesiapan Belajar Siswa pada Materi Rasio"
    )
    
    # Bagian A: Identitas Narasumber
    add_section_heading(doc, "A. Identitas Narasumber / Informan")
    tbl_id = doc.add_table(rows=6, cols=3)
    tbl_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_id, "E2E8F0")
    
    col_widths = [Inches(0.4), Inches(2.3), Inches(3.8)]
    fields = [
        ("1", "Nama Lengkap Guru", ": ....................................................................................."),
        ("2", "NIP / NUPTK", ": ....................................................................................."),
        ("3", "Nama Sekolah / Instansi", ": SMP Negeri 3 Tasikmalaya"),
        ("4", "Pengalaman Mengajar Matematika", ": ........... Tahun"),
        ("5", "Kelas / Fase yang Diampu", ": Kelas VII / Fase D"),
        ("6", "Hari, Tanggal Wawancara", ": .....................................................................................")
    ]
    for r_idx, (no, label, val) in enumerate(fields):
        row = tbl_id.rows[r_idx]
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = w
        format_cell(row.cells[0], no, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[1], label, bold=True, size=11)
        format_cell(row.cells[2], val, size=11)
    
    # Bagian B: Petunjuk Wawancara
    add_section_heading(doc, "B. Petunjuk Pelaksanaan Wawancara")
    add_body_p(doc, "Pedoman wawancara ini digunakan sebagai instrumen studi pendahuluan (tahap Analysis dalam model ADDIE) untuk menggali kondisi faktual, permasalahan pembelajaran matematika di kelas VII, dan kebutuhan guru terhadap inovasi lembar kerja berbantuan kecerdasan artifisial.", "1. Tujuan: ")
    add_body_p(doc, "Wawancara bersifat semiterstruktur. Pewawancara dapat mengajukan pertanyaan pendalaman (probing) secara luwes berdasarkan respon informan, sepanjang tetap berada dalam koridor indikator kebutuhan yang telah ditetapkan.", "2. Sifat: ")
    add_body_p(doc, "Seluruh informasi dan data yang dihimpun semata-mata digunakan untuk kepentingan ilmiah pengembangan skripsi serta dilindungi kerahasiaannya sesuai etika penelitian akademik.", "3. Etika & Privasi: ")
    
    # Bagian C: Kisi-Kisi Pedoman Wawancara
    add_section_heading(doc, "C. Kisi-Kisi Pedoman Wawancara")
    tbl_kisi = doc.add_table(rows=8, cols=5)
    tbl_kisi.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_kisi, "94A3B8")
    
    kisi_widths = [Inches(0.4), Inches(1.8), Inches(3.1), Inches(0.8), Inches(0.6)]
    headers = ["No.", "Aspek Kebutuhan", "Indikator Fokus Penggalian", "Nomor Butir", "Jumlah"]
    for c_idx, h in enumerate(headers):
        tbl_kisi.rows[0].cells[c_idx].width = kisi_widths[c_idx]
        format_cell(tbl_kisi.rows[0].cells[c_idx], h, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color="F1F5F9")
    
    kisi_data = [
        ("1", "Praktik Asesmen", "Bentuk asesmen diagnostik, pemetaan indikator materi rasio, dan tindak lanjut hasil tes", "1, 2, 3", "3"),
        ("2", "Diferensiasi Pembelajaran", "Dasar pengelompokan siswa, adaptasi proses dan konten belajar, hambatan lapangan", "4, 5, 6", "3"),
        ("3", "Penyusunan LKPD", "Alur kerja penyusunan, beban alokasi waktu, komponen lembar kerja, proses revisi", "7, 8, 9", "3"),
        ("4", "Teknologi & AI", "Perangkat lunak yang dipakai, pengalaman pemanfaatan GenAI, harapan fitur sistem", "10, 11, 12", "3"),
        ("5", "Ekspor Formula Matematika", "Kebutuhan format Word (.docx), fasilitas edit rumus OMML, kendala format gambar", "13, 14", "2"),
        ("6", "Privasi & Kontrol Guru", "Minimalisasi data pribadi siswa, peran tinjauan guru (human-in-the-loop)", "15, 16", "2"),
        ("", "Total Butir Pertanyaan", "", "", "16")
    ]
    for r_idx, row_vals in enumerate(kisi_data, start=1):
        row = tbl_kisi.rows[r_idx]
        is_total = (r_idx == 7)
        for c_idx, val in enumerate(row_vals):
            row.cells[c_idx].width = kisi_widths[c_idx]
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
            color = "F8FAFC" if is_total else None
            format_cell(row.cells[c_idx], val, bold=is_total, size=11, align=align, color=color)
    
    # Bagian D: Daftar Butir Pertanyaan Wawancara
    add_section_heading(doc, "D. Daftar Pertanyaan Wawancara Mendalam")
    
    questions = [
        ("I. Aspek Praktik Asesmen Diagnostik", [
            ("1", "Bagaimanakah pelaksanaan asesmen diagnostik kognitif yang selama ini Bapak/Ibu terapkan pada awal pembelajaran materi Rasio di kelas VII? Apakah menggunakan instrumen khusus atau tes formatif biasa?"),
            ("2", "Sejauh mana hasil asesmen diagnostik tersebut dapat memetakan kelemahan spesifik siswa per indikator kompetensi (misalnya membedakan rasio senilai dan berbalik nilai)?"),
            ("3", "Tindak lanjut pedagogis apa yang biasanya Bapak/Ibu lakukan di kelas setelah mengetahui adanya perbedaan tingkat kesiapan belajar siswa yang cukup mencolok?")
        ]),
        ("II. Aspek Pembelajaran Berdiferensiasi dan TaRL", [
            ("4", "Bagaimana cara Bapak/Ibu mengelompokkan siswa berdasarkan tingkat kesiapan belajarnya dalam pembelajaran matematika berdiferensiasi? Apa tantangan utamanya?"),
            ("5", "Bagaimanakah Bapak/Ibu membedakan proses bimbingan (scaffolding) dan konten tugas antara siswa yang masih membutuhkan bimbingan dasar dengan siswa yang sudah mahir?"),
            ("6", "Hambatan praktis apa yang paling sering Bapak/Ibu hadapi saat mencoba menerapkan pendekatan Teaching at the Right Level (TaRL) di tengah keterbatasan jam pelajaran?")
        ]),
        ("III. Aspek Penyusunan dan Penggunaan LKPD", [
            ("7", "Bagaimanakah alur kerja Bapak/Ibu dalam menyusun atau memilih Lembar Kerja Peserta Didik (LKPD) untuk materi Rasio saat ini? Dari mana sumber bahan ajar tersebut diperoleh?"),
            ("8", "Berapa lama rata-rata waktu yang Bapak/Ibu butuhkan jika harus menyusun beberapa variasi LKPD yang disesuaikan dengan tingkat kemampuan siswa yang berbeda-beda?"),
            ("9", "Komponen atau bagian apa saja yang menurut Bapak/Ibu wajib ada dalam LKPD matematika agar mampu memandu alur berpikir siswa secara bertahap (scaffolding)?")
        ]),
        ("IV. Aspek Pemanfaatan Teknologi dan AI", [
            ("10", "Perangkat lunak atau platform digital apa saja yang pernah Bapak/Ibu manfaatkan untuk membantu pembuatan asesmen atau lembar kerja matematika?"),
            ("11", "Apakah Bapak/Ibu pernah mencoba memanfaatkan teknologi Generative AI (seperti ChatGPT atau model kecerdasan artifisial lainnya) dalam merancang bahan ajar? Bagaimana pengalaman dan kendalanya?"),
            ("12", "Fitur otomasi seperti apa yang paling Bapak/Ibu harapkan hadir dalam suatu web application untuk membantu persiapan administrasi diferensiasi pembelajaran?")
        ]),
        ("V. Aspek Kebutuhan Ekspor Formula Matematika", [
            ("13", "Apakah Bapak/Ibu membutuhkan dokumen keluaran berupa file Microsoft Word (.docx) yang dapat diedit bebas kembali sebelum dicetak dan dibagikan ke siswa? Mengapa?"),
            ("14", "Bagaimana kendala yang sering Bapak/Ibu alami terkait penulisan rumus, pecahan, atau simbol matematika ketika mengunduh materi dari internet (misalnya rumus berbentuk gambar yang pecah atau tidak bisa diedit)?")
        ]),
        ("VI. Aspek Privasi Data dan Kontrol Guru", [
            ("15", "Seberapa penting perlindungan data pribadi peserta didik (seperti kerahasiaan nama dan nomor induk) dalam penerapan aplikasi pembelajaran berbasis web di sekolah Bapak/Ibu?"),
            ("16", "Bagaimanakah pandangan Bapak/Ibu mengenai peran guru sebagai peninjau utama (human-in-the-loop) yang berhak mengedit, menyetujui, atau menolak draf materi keluaran AI sebelum digunakan di kelas?")
        ])
    ]
    
    for section_title, q_list in questions:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(2)
        r_sec = p_sec.add_run(section_title)
        r_sec.font.name = "Times New Roman"
        r_sec.font.size = Pt(12)
        r_sec.font.bold = True
        
        for q_num, q_text in q_list:
            p_q = doc.add_paragraph()
            p_q.paragraph_format.line_spacing = 1.15
            p_q.paragraph_format.space_after = Pt(2)
            r_num = p_q.add_run(f"{q_num}. ")
            r_num.font.name = "Times New Roman"
            r_num.font.size = Pt(12)
            r_num.font.bold = True
            r_txt = p_q.add_run(q_text)
            r_txt.font.name = "Times New Roman"
            r_txt.font.size = Pt(12)
            
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.line_spacing = 1.15
            p_ans.paragraph_format.space_after = Pt(6)
            r_ans_lbl = p_ans.add_run("   Catatan Respon Informan:\n")
            r_ans_lbl.font.name = "Times New Roman"
            r_ans_lbl.font.size = Pt(11)
            r_ans_lbl.font.italic = True
            r_dots = p_ans.add_run("   ............................................................................................................................................................\n   ............................................................................................................................................................")
            r_dots.font.name = "Times New Roman"
            r_dots.font.size = Pt(11)
            r_dots.font.color.rgb = RGBColor(148, 163, 184)
            
    # Bagian E: Pengesahan
    add_section_heading(doc, "E. Lembar Konfirmasi dan Catatan Lapangan")
    add_body_p(doc, "Wawancara telah dilaksanakan secara faktual sesuai dengan butir pertanyaan di atas. Data hasil wawancara ini akan dianalisis secara kualitatif sebagai landasan perancangan arsitektur sistem dan konten pembelajaran.")
    
    tbl_ttd = doc.add_table(rows=3, cols=2)
    tbl_ttd.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_ttd, "FFFFFF")
    for r in tbl_ttd.rows:
        r.cells[0].width = Inches(3.2)
        r.cells[1].width = Inches(3.2)
    
    format_cell(tbl_ttd.rows[0].cells[0], "Mengetahui / Menyetujui,\nNarasumber / Guru Matematika,", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(tbl_ttd.rows[0].cells[1], "Tasikmalaya, .................................... 2026\nPewawancara / Peneliti,", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_ttd.rows[1].cells[0], "\n\n\n", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(tbl_ttd.rows[1].cells[1], "\n\n\n", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_ttd.rows[2].cells[0], "( .............................................................. )\nNIP. .....................................................", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(tbl_ttd.rows[2].cells[1], "Ryan Wardiana\nNIM 232151098", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    out_docx = os.path.join(OUT_DIR, "01_Pedoman_Wawancara_Analisis_Kebutuhan_Guru.docx")
    doc.save(out_docx)
    print("Saved:", out_docx)

if __name__ == "__main__":
    build_instrument_01()

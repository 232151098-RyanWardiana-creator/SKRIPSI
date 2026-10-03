import os, sys
sys.path.append('.')
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from scripts.generate_all_instruments import (
    create_base_doc, add_header_kop, add_doc_title, add_section_heading,
    add_body_p, format_cell, set_table_borders, OUT_DIR
)

def build_instrument_03():
    doc = create_base_doc()
    add_header_kop(doc)
    add_doc_title(
        doc,
        "LEMBAR VALIDASI AHLI MEDIA DAN TEKNOLOGI PENDIDIKAN",
        "Penilaian Kelayakan Teknis Web Application Berbasis Artificial Intelligence untuk Pembelajaran Matematika Berdiferensiasi"
    )
    
    # Bagian A: Identitas Validator
    add_section_heading(doc, "A. Identitas Validator Ahli Media")
    tbl_id = doc.add_table(rows=5, cols=3)
    tbl_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_id, "E2E8F0")
    col_widths = [Inches(0.4), Inches(2.3), Inches(3.8)]
    fields = [
        ("1", "Nama Lengkap & Gelar", ": ....................................................................................."),
        ("2", "NIP / NIDN", ": ....................................................................................."),
        ("3", "Instansi / Perguruan Tinggi", ": ....................................................................................."),
        ("4", "Keahlian / Bidang Minat", ": Teknologi Pendidikan / Rekayasa Perangkat Lunak / UI-UX"),
        ("5", "Hari, Tanggal Penilaian", ": .....................................................................................")
    ]
    for r_idx, (no, label, val) in enumerate(fields):
        row = tbl_id.rows[r_idx]
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = w
        format_cell(row.cells[0], no, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[1], label, bold=True, size=11)
        format_cell(row.cells[2], val, size=11)
    
    # Bagian B: Petunjuk Penilaian
    add_section_heading(doc, "B. Petunjuk Pengisian Lembar Validasi")
    add_body_p(doc, "Lembar validasi ini bertujuan untuk menghimpun data penilaian para ahli mengenai kualitas teknis, desain antarmuka (UI/UX), keandalan sistem kecerdasan artifisial, dan integritas keluaran berkas dokumen web application yang dikembangkan.", "1. Tujuan: ")
    add_body_p(doc, "Bapak/Ibu dimohon memberikan tanda centang (✓) pada kolom skala nilai (1–5) yang paling mencerminkan penilaian objektif Bapak/Ibu terhadap setiap butir pernyataan.", "2. Cara Penilaian: ")
    add_body_p(doc, "Kriteria skala penilaian: 5 = Sangat Baik (SB); 4 = Baik (B); 3 = Cukup (C); 2 = Kurang (K); 1 = Sangat Kurang (SK).", "3. Kategori Skala: ")
    add_body_p(doc, "Bapak/Ibu dimohon memberikan masukan, catatan teknis, atau rekomendasi penyempurnaan pada kolom yang disediakan.", "4. Saran Perbaikan: ")
    
    # Bagian C: Tabel Penilaian 40 Butir (Tabel 7.4)
    add_section_heading(doc, "C. Penilaian Kelayakan Media dan Teknis Perangkat Lunak")
    
    items_data = [
        # Aspek 1: Fungsionalitas Aplikasi (1-7)
        ("I. Aspek Fungsionalitas Aplikasi", [
            ("1", "Kemudahan alur akses masuk (login/onboarding) pengguna ke dalam lingkungan web application."),
            ("2", "Keberfungsian modul formulir pelaksanaan tes asesmen diagnostik daring bagi peserta didik."),
            ("3", "Keberfungsian sistem otomasi penskoran tes diagnostik berdasarkan kunci jawaban dan rubrik."),
            ("4", "Keberfungsian algoritma pemetaan otomatis kesiapan belajar ke dalam tiga kategori kurikuler."),
            ("5", "Keberfungsian modul generator AI dalam memproduksi tiga draf varian LKPD secara simultan."),
            ("6", "Keberfungsian editor teks interaktif untuk pratinjau dan penyuntingan konten oleh guru."),
            ("7", "Keberfungsian seluruh tombol aksi, alur navigasi, dan transisi halaman tanpa galat (bug).")
        ]),
        # Aspek 2: Antarmuka Pengguna (UI) dan Navigasi (8-13)
        ("II. Aspek Antarmuka Pengguna (UI) dan Navigasi", [
            ("8", "Konsistensi tata letak (layout), skema warna, dan hierarki visual antarhalaman aplikasi."),
            ("9", "Kejelasan penataan menu navigasi utama, tombol fungsi, dan bilah progres pengerjaan."),
            ("10", "Daya tarik visual antarmuka sistem yang bersih, modern, profesional, dan ramah pengguna."),
            ("11", "Kemudahan mengenali simbol, ikonografi, dan status penanda proses pada aplikasi."),
            ("12", "Ketersediaan umpan balik visual langsung (visual feedback/toast alert) atas aksi pengguna."),
            ("13", "Kenyamanan tata letak komponen yang tidak menimbulkan kelelahan visual (eye strain).")
        ]),
        # Aspek 3: Keterbacaan dan Responsivitas (14-17)
        ("III. Aspek Keterbacaan dan Responsivitas", [
            ("14", "Keterbacaan tipografi (pemilihan jenis font, ukuran, dan jarak spasi baris) pada antarmuka."),
            ("15", "Kontras rasio warna yang memadai antara warna teks dan latar belakang sesuai kaidah desain."),
            ("16", "Ketepatan rendering formula, pecahan, dan simbol matematika di layar peramban tanpa terpotong."),
            ("17", "Responsivitas tampilan sistem yang adaptif saat diakses via layar komputer, laptop, maupun tablet.")
        ]),
        # Aspek 4: Aksesibilitas Sistem (18-21)
        ("IV. Aspek Aksesibilitas Sistem", [
            ("18", "Kejelasan label, teks petunjuk, dan deskripsi pada setiap kolom formulir masukan data."),
            ("19", "Kemudahan aksesibilitas navigasi aplikasi menggunakan papan ketik (keyboard accessibility)."),
            ("20", "Kejelasan, ketepatan, dan keinformatifan pesan kesalahan (error handling message) kepada pengguna."),
            ("21", "Ketersediaan dokumentasi bantuan atau panduan ringkas penggunaan sistem yang mudah dipahami.")
        ]),
        # Aspek 5: Reliabilitas Sistem (22-25)
        ("V. Aspek Reliabilitas dan Penanganan Kesalahan", [
            ("22", "Efektivitas validasi masukan data formulir untuk mencegah kesalahan fatal akibat input pengguna."),
            ("23", "Ketahanan sistem terhadap gangguan fluktuasi jaringan internet atau latensi koneksi."),
            ("24", "Ketersediaan proteksi pencegahan kehilangan data draf lembar kerja guru yang belum tersimpan."),
            ("25", "Stabilitas performa aplikasi saat mengeksekusi proses komputasi yang berulang.")
        ]),
        # Aspek 6: Privasi dan Keamanan Data (26-29)
        ("VI. Aspek Privasi dan Keamanan Data", [
            ("26", "Penerapan prinsip minimalisasi data pribadi dalam perekaman identitas peserta didik."),
            ("27", "Keamanan mekanisme otentikasi sesi guru dan isolasi data antarakun pengguna."),
            ("28", "Perlindungan integritas dan kerahasiaan basis data hasil tes diagnostik peserta didik."),
            ("29", "Keamanan transmisi data antara front-end, back-end server, dan penyedia API kecerdasan artifisial.")
        ]),
        # Aspek 7: Alur Kerja GenAI dan Kontrol Guru (30-33)
        ("VII. Aspek Alur Kerja GenAI dan Kontrol Guru", [
            ("30", "Keinformatifan indikator proses generasi teks AI (loading indicator/progress state) bagi pengguna."),
            ("31", "Keterlacakan dan transparansi parameter prompt sistem dalam mengarahkan pembentukan bahan ajar."),
            ("32", "Kemudahan fasilitas interaktif bagi guru untuk membatalkan, merestart, atau merevisi draf AI."),
            ("33", "Penempatan teknologi AI secara etis dan proporsional sebagai asisten penguat (bukan penentu mutlak).")
        ]),
        # Aspek 8: Ekspor Dokumen Microsoft Word (.docx) (34-38)
        ("VIII. Aspek Kualitas Ekspor Dokumen Word (.docx)", [
            ("34", "Integritas berkas .docx hasil unduhan (dapat dibuka lancar di MS Word tanpa peringatan korup)."),
            ("35", "Keberhasilan konversi formula matematika web menjadi formula asli Microsoft Word (OMML)."),
            ("36", "Kemampuan rumus matematika pada dokumen Word hasil unduhan untuk disunting langsung (editable)."),
            ("37", "Kerapian tata letak dokumen Word (margin, tabel, ruang jawaban peserta didik, dan penomoran)."),
            ("38", "Kesesuaian format tata letak berkas Word dengan standar cetak kertas A4 siap pakai di kelas.")
        ]),
        # Aspek 9: Kinerja dan Efisiensi Teknis (39-40)
        ("IX. Aspek Kinerja dan Efisiensi Teknis", [
            ("39", "Waktu respons server dan kecepatan rendering halaman antarmuka dalam batas toleransi wajar."),
            ("40", "Efisiensi waktu proses komputasi generasi draf, konversi formula, dan pengunduhan berkas .docx.")
        ])
    ]
    
    tbl_eval = doc.add_table(rows=1, cols=8)
    tbl_eval.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_eval, "94A3B8")
    col_w = [Inches(0.4), Inches(3.4), Inches(0.4), Inches(0.4), Inches(0.4), Inches(0.4), Inches(0.4), Inches(1.3)]
    headers = ["No.", "Indikator Penilaian Media & Teknologi", "1", "2", "3", "4", "5", "Catatan / Saran"]
    
    for c_idx, h in enumerate(headers):
        tbl_eval.rows[0].cells[c_idx].width = col_w[c_idx]
        format_cell(tbl_eval.rows[0].cells[c_idx], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color="F1F5F9")
    
    for group_title, items in items_data:
        hdr_row = tbl_eval.add_row()
        for c in hdr_row.cells:
            for w_idx, w in enumerate(col_w):
                hdr_row.cells[w_idx].width = w
        format_cell(hdr_row.cells[0], "", bold=True, size=10, color="E2E8F0")
        format_cell(hdr_row.cells[1], group_title, bold=True, size=10, color="E2E8F0")
        for c_idx in range(2, 8):
            format_cell(hdr_row.cells[c_idx], "", color="E2E8F0")
            
        for num, text in items:
            r = tbl_eval.add_row()
            for w_idx, w in enumerate(col_w):
                r.cells[w_idx].width = w
            format_cell(r.cells[0], num, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_cell(r.cells[1], text, size=10)
            for c_idx in range(2, 7):
                format_cell(r.cells[c_idx], "[  ]", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_cell(r.cells[7], "", size=9)
            
    # Bagian D: Catatan dan Saran Umum
    add_section_heading(doc, "D. Catatan Teknis dan Saran Perbaikan Umum")
    tbl_notes = doc.add_table(rows=1, cols=1)
    tbl_notes.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_notes.rows[0].cells[0].width = Inches(7.1)
    format_cell(tbl_notes.rows[0].cells[0], "\n\n\n\n\n\n", size=11)
    set_table_borders(tbl_notes, "94A3B8")
    
    # Bagian E: Kesimpulan Rekomendasi
    add_section_heading(doc, "E. Rekomendasi Kelayakan Media")
    add_body_p(doc, "Berdasarkan evaluasi menyeluruh terhadap aspek fungsionalitas, antarmuka, aksesibilitas, reliabilitas, keamanan data, performa AI, dan ekspor berkas Word, web application ini dinyatakan:")
    add_body_p(doc, "Layak diujicobakan di sekolah tanpa perbaikan sistem.", "[   ] A. ")
    add_body_p(doc, "Layak diujicobakan dengan perbaikan teknis minor sesuai catatan.", "[   ] B. ")
    add_body_p(doc, "Belum layak diujicobakan / memerlukan perbaikan arsitektur mayor.", "[   ] C. ")
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(8)
    
    tbl_sign = doc.add_table(rows=3, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sign, "FFFFFF")
    for r in tbl_sign.rows:
        r.cells[0].width = Inches(3.5)
        r.cells[1].width = Inches(3.5)
        
    format_cell(tbl_sign.rows[0].cells[0], "", size=11)
    format_cell(tbl_sign.rows[0].cells[1], "Tasikmalaya, .................................... 2026\nValidator Ahli Media,", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_sign.rows[1].cells[0], "", size=11)
    format_cell(tbl_sign.rows[1].cells[1], "\n\n\n", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_sign.rows[2].cells[0], "", size=11)
    format_cell(tbl_sign.rows[2].cells[1], "( .............................................................. )\nNIP/NIDN. ...............................................", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    out_docx = os.path.join(OUT_DIR, "03_Lembar_Validasi_Ahli_Media.docx")
    doc.save(out_docx)
    print("Saved:", out_docx)

if __name__ == "__main__":
    build_instrument_03()

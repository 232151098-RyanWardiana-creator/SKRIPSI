import os, sys
sys.path.append('.')
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from scripts.generate_all_instruments import (
    create_base_doc, add_header_kop, add_doc_title, add_section_heading,
    add_body_p, format_cell, set_table_borders, OUT_DIR
)

def build_instrument_05():
    doc = create_base_doc()
    add_header_kop(doc)
    add_doc_title(
        doc,
        "LEMBAR VALIDASI ISI BUTIR SOAL TES ASESMEN DIAGNOSTIK",
        "Pengujian Validitas Isi Berdasarkan Koefisien Aiken's V oleh Panel Ahli Pembelajaran Matematika"
    )
    
    # Bagian A: Identitas Validator
    add_section_heading(doc, "A. Identitas Validator / Penilai")
    tbl_id = doc.add_table(rows=5, cols=3)
    tbl_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_id, "E2E8F0")
    col_widths = [Inches(0.4), Inches(2.3), Inches(3.8)]
    fields = [
        ("1", "Nama Lengkap & Gelar", ": ....................................................................................."),
        ("2", "NIP / NIDN", ": ....................................................................................."),
        ("3", "Instansi / Perguruan Tinggi", ": ....................................................................................."),
        ("4", "Keahlian / Bidang Minat", ": Evaluasi Pembelajaran Matematika / Pendidikan Matematika"),
        ("5", "Hari, Tanggal Penilaian", ": .....................................................................................")
    ]
    for r_idx, (no, label, val) in enumerate(fields):
        row = tbl_id.rows[r_idx]
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = w
        format_cell(row.cells[0], no, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[1], label, bold=True, size=11)
        format_cell(row.cells[2], val, size=11)
        
    # Bagian B: Petunjuk Penilaian & Formula Aiken
    add_section_heading(doc, "B. Petunjuk Penilaian dan Kriteria Relevansi")
    add_body_p(doc, "Lembar validasi ini digunakan oleh panel ahli untuk menelaah kesesuaian, keterwakilan, dan relevansi butir instrumen tes asesmen diagnostik terhadap lima indikator ketercapaian kompetensi (IK-01 s.d. IK-05) materi rasio Fase D.", "1. Tujuan: ")
    add_body_p(doc, "Bapak/Ibu dimohon menelaah butir soal ditinjau dari aspek: (a) Substansi materi, (b) Konstruksi soal, (c) Keterbacaan bahasa, dan (d) Kebenaran kunci jawaban serta rubrik.", "2. Aspek Telaah: ")
    add_body_p(doc, "Berikan tanda centang (✓) pada kolom derajat relevansi (1–5) dengan ketentuan: 1 = Sangat Tidak Relevan (STR); 2 = Kurang Relevan (KR); 3 = Cukup Relevan (CR); 4 = Relevan (R); 5 = Sangat Relevan (SR).", "3. Skala Penilaian: ")
    add_body_p(doc, "Data skor relevansi dianalisis menggunakan koefisien V Aiken: V = Σs / [n × (c - 1)], di mana s = r - l₀, n = jumlah validator, dan c = skor tertinggi (5). Butir soal dinyatakan valid jika koefisien V ≥ 0,80.", "4. Analisis Aiken's V: ")

    # Bagian C: Tabel Penilaian Validitas Isi (10 Butir)
    add_section_heading(doc, "C. Lembar Penilaian Derajat Relevansi Butir Soal")
    
    tbl_aiken = doc.add_table(rows=1, cols=9)
    tbl_aiken.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_aiken, "94A3B8")
    col_w = [Inches(0.4), Inches(0.8), Inches(2.6), Inches(0.4), Inches(0.4), Inches(0.4), Inches(0.4), Inches(0.4), Inches(1.3)]
    headers = ["No.", "Kode IK", "Deskripsi Butir & Aspek Ukur", "1", "2", "3", "4", "5", "Saran Perbaikan Butir"]
    
    for c_idx, h in enumerate(headers):
        tbl_aiken.rows[0].cells[c_idx].width = col_w[c_idx]
        format_cell(tbl_aiken.rows[0].cells[c_idx], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color="F1F5F9")
        
    items_aiken = [
        ("1", "IK-01", "Soal No. 1: Mengukur kemampuan menyatakan rasio dua himpunan objek sejenis (siswa laki-laki vs perempuan) ke bentuk rasio paling sederhana dan pecahan."),
        ("2", "IK-01", "Soal No. 2: Mengukur kemampuan membandingkan dua besaran dengan satuan berbeda (meter vs sentimeter) melalui konversi penyetaraan satuan."),
        ("3", "IK-02", "Soal No. 3: Mengukur keterampilan menyederhanakan rasio campuran bahan minuman dengan menentukan faktor pembagi terbesar (FPB)."),
        ("4", "IK-02", "Soal No. 4: Mengukur kemampuan menentukan rasio ekuivalen (senilai) untuk mencari nilai kuantitas yang belum diketahui dalam perbandingan kelereng."),
        ("5", "IK-03", "Soal No. 5: Mengukur kemampuan menentukan rasio satuan laju kecepatan (km/jam) untuk membandingkan kelajuan dua mobil secara objektif."),
        ("6", "IK-03", "Soal No. 6: Mengukur kemampuan menentukan rasio harga satuan (rupiah/kg) untuk mengambil keputusan pembelian kemasan beras yang lebih ekonomis."),
        ("7", "IK-04", "Soal No. 7: Mengukur penalaran pemecahan masalah kontekstual perbandingan senilai (kebutuhan bahan bakar terhadap jarak tempuh kendaraan)."),
        ("8", "IK-04", "Soal No. 8: Mengukur penalaran pemecahan masalah kontekstual perbandingan berbalik nilai (percepatan waktu proyek terhadap kebutuhan pekerja tambahan)."),
        ("9", "IK-05", "Soal No. 9 (HOTS): Mengukur kemampuan bernalar tinggi dalam menyelesaikan masalah skala peta provinsi, konversi jarak nyata, dan estimasi waktu perjalanan."),
        ("10", "IK-05", "Soal No. 10 (HOTS): Mengukur kemampuan pemecahan masalah multidimensi (denah skala, luas ruang laboratorium sebenarnya, dan kapasitas daya tampung meja komputer).")
    ]
    
    for no, ik, desc in items_aiken:
        r = tbl_aiken.add_row()
        for c_idx, w in enumerate(col_w):
            r.cells[c_idx].width = w
        format_cell(r.cells[0], no, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(r.cells[1], ik, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(r.cells[2], desc, size=9.5)
        for c_idx in range(3, 8):
            format_cell(r.cells[c_idx], "[  ]", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(r.cells[8], "", size=9)
        
    # Bagian D: Catatan Umum
    add_section_heading(doc, "D. Catatan Kualitatif dan Saran Penyempurnaan Soal")
    tbl_notes = doc.add_table(rows=1, cols=1)
    tbl_notes.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_notes.rows[0].cells[0].width = Inches(7.1)
    format_cell(tbl_notes.rows[0].cells[0], "\n\n\n\n\n", size=11)
    set_table_borders(tbl_notes, "94A3B8")
    
    # Bagian E: Kesimpulan Rekomendasi
    add_section_heading(doc, "E. Rekomendasi Panel Validator")
    add_body_p(doc, "Berdasarkan telaah butir soal asesmen diagnostik terhadap kisi-kisi dan indikator kompetensi, instrumen tes ini dinyatakan:")
    add_body_p(doc, "Dapat digunakan untuk asesmen diagnostik tanpa revisi.", "[   ] A. ")
    add_body_p(doc, "Dapat digunakan untuk asesmen diagnostik dengan revisi sesuai catatan.", "[   ] B. ")
    add_body_p(doc, "Tidak dapat digunakan / butir soal harus diganti.", "[   ] C. ")
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(8)
    
    tbl_sign = doc.add_table(rows=3, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sign, "FFFFFF")
    for r in tbl_sign.rows:
        r.cells[0].width = Inches(3.5)
        r.cells[1].width = Inches(3.5)
        
    format_cell(tbl_sign.rows[0].cells[0], "", size=11)
    format_cell(tbl_sign.rows[0].cells[1], "Tasikmalaya, .................................... 2026\nValidator Ahli Evaluasi,", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_sign.rows[1].cells[0], "", size=11)
    format_cell(tbl_sign.rows[1].cells[1], "\n\n\n", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_sign.rows[2].cells[0], "", size=11)
    format_cell(tbl_sign.rows[2].cells[1], "( .............................................................. )\nNIP/NIDN. ...............................................", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    out_docx = os.path.join(OUT_DIR, "05_Lembar_Validasi_Isi_Tes_Diagnostik_Aikens_V.docx")
    doc.save(out_docx)
    print("Saved:", out_docx)

if __name__ == "__main__":
    build_instrument_05()

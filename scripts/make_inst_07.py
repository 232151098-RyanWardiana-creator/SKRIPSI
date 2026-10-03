import os, sys
sys.path.append('.')
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from scripts.generate_all_instruments import (
    create_base_doc, add_header_kop, add_doc_title, add_section_heading,
    add_body_p, format_cell, set_table_borders, OUT_DIR
)

def build_instrument_07():
    doc = create_base_doc()
    add_header_kop(doc)
    add_doc_title(
        doc,
        "ANGKET RESPON KEPRAKTISAN BELAJAR PESERTA DIDIK",
        "Uji Coba Lapangan E-LKPD Matematika Berdiferensiasi Berbasis Web pada Materi Rasio Kelas VII SMP"
    )
    
    # Bagian A: Identitas Siswa
    add_section_heading(doc, "A. Identitas Peserta Didik")
    tbl_id = doc.add_table(rows=5, cols=3)
    tbl_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_id, "E2E8F0")
    col_widths = [Inches(0.4), Inches(2.0), Inches(4.1)]
    fields = [
        ("1", "Nama Lengkap / Inisial", ": ....................................................................................."),
        ("2", "Nomor Induk / Absen", ": ....................................................................................."),
        ("3", "Kelas / Sekolah", ": Kelas VII-..... / SMP Negeri 3 Tasikmalaya"),
        ("4", "Jenis Kelamin", ": [   ] Laki-laki        [   ] Perempuan"),
        ("5", "Hari, Tanggal Pengisian", ": .....................................................................................")
    ]
    for r_idx, (no, label, val) in enumerate(fields):
        row = tbl_id.rows[r_idx]
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = w
        format_cell(row.cells[0], no, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[1], label, bold=True, size=11)
        format_cell(row.cells[2], val, size=11)
        
    # Bagian B: Petunjuk Pengisian
    add_section_heading(doc, "B. Petunjuk Pengisian Angket Respon")
    add_body_p(doc, "Angket ini bertujuan untuk mengetahui pengalaman, kemudahan, dan perasaanmu saat belajar matematika materi rasio menggunakan lembar kerja (LKPD) yang baru saja kamu kerjakan.", "1. Halo Adik-adik: ")
    add_body_p(doc, "Jawaban yang kamu berikan tidak akan mempengaruhi nilai rapotmu sama sekali, jadi jawablah secara jujur sesuai apa yang benar-benar kamu rasakan.", "2. Tidak Mempengaruhi Nilai: ")
    add_body_p(doc, "Berikan tanda centang (✓) pada kotak pilihan yang paling cocok: STS = Sangat Tidak Setuju; TS = Tidak Setuju; CS = Cukup Setuju / Netral; S = Setuju; SS = Sangat Setuju.", "3. Cara Menjawab: ")

    # Bagian C: Tabel 25 Butir (Tabel 7.6)
    add_section_heading(doc, "C. Lembar Respon Kepraktisan Peserta Didik")
    
    items_data = [
        # Aspek 1: Akses & Tampilan Awal (1-3)
        ("I. Aspek Kemudahan Akses dan Tampilan Web", [
            ("1", "Saya merasa mudah membuka dan masuk ke dalam tautan web pembelajaran matematika ini."),
            ("2", "Tampilan awal web terlihat bersih, menarik, dan tidak membuat saya merasa bingung."),
            ("3", "Saya dapat langsung menemukan menu untuk mulai mengerjakan tes atau membuka lembar kerja.")
        ]),
        # Aspek 2: Kejelasan Petunjuk (4-7)
        ("II. Aspek Kejelasan Petunjuk Belajar", [
            ("4", "Petunjuk cara mengerjakan tes diagnostik tertulis dengan kalimat yang sangat jelas dan mudah dipahami."),
            ("5", "Langkah-langkah kegiatan belajar pada lembar kerja (LKPD) mudah saya ikuti dari awal hingga akhir."),
            ("6", "Saya selalu tahu apa yang harus saya lakukan saat menyelesaikan setiap soal atau tugas yang ada."),
            ("7", "Tanda atau petunjuk bantuan pada lembar kerja mudah dikenali dan sangat menuntun saya.")
        ]),
        # Aspek 3: Keterbacaan Teks, Gambar, dan Rumus (8-11)
        ("III. Aspek Keterbacaan Teks, Gambar, dan Rumus", [
            ("8", "Ukuran tulisan pada lembar kerja pas, jelas, dan sangat nyaman dibaca oleh mata saya."),
            ("9", "Gambar dan ilustrasi cerita yang disajikan terlihat jelas dan membantu saya membayangkan masalahnya."),
            ("10", "Tulisan angka, pecahan, dan rumus matematika tersusun rapi serta tidak membingungkan saya."),
            ("11", "Bahasa yang digunakan dalam lembar kerja terasa akrab, ramah, dan mudah saya mengerti.")
        ]),
        # Aspek 4: Interaksi & Pengerjaan (12-15)
        ("IV. Aspek Kemudahan Interaksi dan Pengerjaan Soal", [
            ("12", "Saya dapat mengisi atau menuliskan jawaban latihan dengan lancar tanpa ada kesulitan teknis."),
            ("13", "Berpindah dari satu halaman lembar kerja ke halaman berikutnya terasa sangat mudah."),
            ("14", "Tempat yang disediakan untuk menuliskan cara coretan hitung dan jawaban akhir sangat cukup."),
            ("15", "Saya tidak merasa pusing saat menggunakan perangkat ini selama jam pelajaran matematika.")
        ]),
        # Aspek 5: Bantuan Belajar / Scaffolding (16-18)
        ("V. Aspek Kesesuaian Bantuan Belajar (Scaffolding)", [
            ("16", "Contoh penyelesaian dan petunjuk langkah awal sangat membantu saya ketika saya merasa ragu."),
            ("17", "Tingkat kesulitan soal pada lembar kerja yang saya terima terasa pas dengan kemampuan saya saat ini."),
            ("18", "Bantuan langkah-langkah belajar membuat saya menjadi lebih berani mencoba menyelesaikan soal sendiri.")
        ]),
        # Aspek 6: Umpan Balik Kemajuan (19-21)
        ("VI. Aspek Umpan Balik dan Informasi Kemajuan", [
            ("19", "Saya dapat mengetahui tugas mana yang sudah selesai dan mana yang masih harus saya kerjakan."),
            ("20", "Informasi hasil tes diagnostik membuat saya tahu bagian materi rasio mana yang sudah saya kuasai."),
            ("21", "Saya merasa lebih bersemangat untuk berlatih lebih giat pada bagian soal yang belum saya pahami.")
        ]),
        # Aspek 7: Kenyamanan & Semangat (22-25)
        ("VII. Aspek Kenyamanan dan Semangat Belajar", [
            ("22", "Saya merasa senang dan lebih bersemangat belajar materi rasio menggunakan lembar kerja ini."),
            ("23", "Saya merasa tidak terbebani karena lembar kerja yang saya terima sesuai dengan kecepatan belajar saya."),
            ("24", "Belajar matematika di kelas menjadi terasa lebih seru, mengasyikkan, dan tidak membosankan."),
            ("25", "Saya berharap guru matematika sering menggunakan lembar kerja model ini pada pelajaran berikutnya.")
        ])
    ]
    
    tbl_eval = doc.add_table(rows=1, cols=7)
    tbl_eval.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_eval, "94A3B8")
    col_w = [Inches(0.4), Inches(4.3), Inches(0.45), Inches(0.45), Inches(0.45), Inches(0.45), Inches(0.45)]
    headers = ["No.", "Pernyataan Respon Peserta Didik", "STS", "TS", "CS", "S", "SS"]
    
    for c_idx, h in enumerate(headers):
        tbl_eval.rows[0].cells[c_idx].width = col_w[c_idx]
        format_cell(tbl_eval.rows[0].cells[c_idx], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color="F1F5F9")
        
    for group_title, items in items_data:
        hdr_row = tbl_eval.add_row()
        for w_idx, w in enumerate(col_w):
            hdr_row.cells[w_idx].width = w
        format_cell(hdr_row.cells[0], "", bold=True, size=10, color="E2E8F0")
        format_cell(hdr_row.cells[1], group_title, bold=True, size=10, color="E2E8F0")
        for c_idx in range(2, 7):
            format_cell(hdr_row.cells[c_idx], "", color="E2E8F0")
            
        for num, text in items:
            r = tbl_eval.add_row()
            for w_idx, w in enumerate(col_w):
                r.cells[w_idx].width = w
            format_cell(r.cells[0], num, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_cell(r.cells[1], text, size=10)
            for c_idx in range(2, 7):
                format_cell(r.cells[c_idx], "[  ]", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                
    # Bagian D: Refleksi Terbuka Siswa
    add_section_heading(doc, "D. Kesan, Cerita, dan Harapan Siswa (Refleksi)")
    add_body_p(doc, "1. Apa hal yang paling kamu sukai saat belajar menggunakan lembar kerja (LKPD) ini?")
    tbl_k1 = doc.add_table(rows=1, cols=1)
    tbl_k1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_k1.rows[0].cells[0].width = Inches(7.1)
    format_cell(tbl_k1.rows[0].cells[0], "\n\n", size=11)
    set_table_borders(tbl_k1, "94A3B8")
    
    add_body_p(doc, "2. Apakah ada bagian materi rasio atau langkah di lembar kerja yang menurutmu masih sulit? Ceritakan!")
    tbl_k2 = doc.add_table(rows=1, cols=1)
    tbl_k2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_k2.rows[0].cells[0].width = Inches(7.1)
    format_cell(tbl_k2.rows[0].cells[0], "\n\n", size=11)
    set_table_borders(tbl_k2, "94A3B8")
    
    # Tanda tangan
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(8)
    
    tbl_sign = doc.add_table(rows=3, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sign, "FFFFFF")
    for r in tbl_sign.rows:
        r.cells[0].width = Inches(3.5)
        r.cells[1].width = Inches(3.5)
        
    format_cell(tbl_sign.rows[0].cells[0], "", size=11)
    format_cell(tbl_sign.rows[0].cells[1], "Tasikmalaya, .................................... 2026\nPeserta Didik,", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_sign.rows[1].cells[0], "", size=11)
    format_cell(tbl_sign.rows[1].cells[1], "\n\n\n", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_sign.rows[2].cells[0], "", size=11)
    format_cell(tbl_sign.rows[2].cells[1], "( .............................................................. )\nNISN. .....................................................", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    out_docx = os.path.join(OUT_DIR, "07_Angket_Kepraktisan_Siswa.docx")
    doc.save(out_docx)
    print("Saved:", out_docx)

if __name__ == "__main__":
    build_instrument_07()

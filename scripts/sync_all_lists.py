import docx

DOCX_PATH = r"C:\Users\LATITUDE 3310 TOUCH\Documents\DOCUMENT RYAN\02_Kuliah\Semester_7\KF21518001 - Skripsi (Kelas A)\03_Draft Proposal (BAB 1-3)\Proposal_Skripsi_Ryan_Wardiana_Revisi_Bimbingan_2_Draft.docx"

def sync_lists():
    doc = docx.Document(DOCX_PATH)
    body = doc._element.body

    # =========================================================================
    # 1. Sinkronisasi SDT (DAFTAR ISI) - Singkatan Rapi 1 Baris
    # =========================================================================
    sdt = next(el for el in body if el.tag.endswith('sdt'))
    paras = list(sdt.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'))

    # Pemetaan indeks SDT ke judul ringkas dan nomor halaman aktual
    toc_updates = {
        9:  ("3.1", "Web Application LKPD Berbasis GenAI", "4"),
        10: ("3.2", "GenAI dan LLM", "4"),
        11: ("3.3", "Formula Matematika dan OMML", "4"),
        12: ("3.4", "Pembelajaran Berdiferensiasi dan TaRL", "5"),
        13: ("3.5", "Asesmen Diagnostik dan Kesiapan Belajar", "5"),
        14: ("3.6", "Scaffolding dan HOTS", "5"),
        15: ("3.7", "Materi Rasio dan Perbandingan", "6"),
        22: ("6.1.1", "E-LKPD Berbasis Web dan ADDIE", "8"),
        23: ("6.1.2", "GenAI dan LLM", "10"),
        24: ("6.1.3", "Formula Matematika dan OMML", "11"),
        25: ("6.1.4", "Pembelajaran Berdiferensiasi dan TaRL", "13"),
        26: ("6.1.5", "Asesmen Diagnostik dan Kesiapan Belajar", "14"),
        27: ("6.1.6", "Scaffolding dan HOTS", "16"),
        28: ("6.1.7", "Materi Rasio dan Perbandingan", "18"),
        41: ("7.5", "Teknik Analisis Data", "32"),
        42: ("7.6", "Waktu dan Tempat Penelitian", "33"),
    }

    for idx, (num, title, pg) in toc_updates.items():
        if idx < len(paras):
            p_el = paras[idx]
            t_nodes = list(p_el.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
            if len(t_nodes) >= 3:
                t_nodes[0].text = num
                t_nodes[1].text = title
                t_nodes[2].text = pg
                print(f"   Updated TOC[{idx:2d}]: {num} {title} -> Hlm {pg}")
            elif len(t_nodes) == 2:
                t_nodes[0].text = title
                t_nodes[1].text = pg
                print(f"   Updated TOC[{idx:2d}]: {title} -> Hlm {pg}")

    # =========================================================================
    # 2. Sinkronisasi DAFTAR TABEL (Tabel 7.1 s.d. 7.8)
    # =========================================================================
    table_data = [
        ("Tabel 7.1 Sumber Data dan Peran Subjek Penelitian", "25"),
        ("Tabel 7.2 Kisi-Kisi Pedoman Wawancara Kebutuhan Guru", "27"),
        ("Tabel 7.3 Kisi-Kisi Lembar Validasi Ahli Materi", "27"),
        ("Tabel 7.4 Kisi-Kisi Lembar Validasi Ahli Media", "29"),
        ("Tabel 7.5 Kisi-Kisi Angket Kepraktisan Guru", "30"),
        ("Tabel 7.6 Kisi-Kisi Angket Kepraktisan Siswa", "30"),
        ("Tabel 7.7 Kriteria Interval Kepraktisan dan Respons Pengguna", "33"),
        ("Tabel 7.8 Matriks Jadwal Rencana Penelitian dan Pengembangan", "35"),
    ]

    # Temukan paragraf "DAFTAR TABEL"
    idx_dt = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "DAFTAR TABEL":
            idx_dt = i
            break

    if idx_dt is not None:
        print(f"\n   Found DAFTAR TABEL at paragraph {idx_dt}")
        # Hapus paragraf lama antara DAFTAR TABEL dan DAFTAR GAMBAR
        curr = idx_dt + 1
        while curr < len(doc.paragraphs) and doc.paragraphs[curr].text.strip() != "DAFTAR GAMBAR":
            p_rem = doc.paragraphs[curr]._p
            p_rem.getparent().remove(p_rem)

        # Sisipkan 8 baris tabel baru yang rapi
        ref_p = doc.paragraphs[idx_dt]._p
        parent = ref_p.getparent()
        ref_idx = parent.index(ref_p)

        for offset, (t_title, t_pg) in enumerate(table_data, start=1):
            new_p = docx.oxml.OxmlElement('w:p')
            new_pPr = docx.oxml.OxmlElement('w:pPr')
            new_pStyle = docx.oxml.OxmlElement('w:pStyle')
            new_pStyle.set(docx.oxml.ns.qn('w:val'), 'TOC1')
            new_pPr.append(new_pStyle)
            new_p.append(new_pPr)

            # Run 1: Judul Tabel
            r1 = docx.oxml.OxmlElement('w:r')
            rPr1 = docx.oxml.OxmlElement('w:rPr')
            rf1 = docx.oxml.OxmlElement('w:rFonts')
            rf1.set(docx.oxml.ns.qn('w:ascii'), 'Times New Roman')
            rf1.set(docx.oxml.ns.qn('w:hAnsi'), 'Times New Roman')
            sz1 = docx.oxml.OxmlElement('w:sz')
            sz1.set(docx.oxml.ns.qn('w:val'), '24')
            rPr1.append(rf1)
            rPr1.append(sz1)
            t1 = docx.oxml.OxmlElement('w:t')
            t1.text = t_title
            r1.append(rPr1)
            r1.append(t1)
            new_p.append(r1)

            # Run 2: Tab Leader
            r2 = docx.oxml.OxmlElement('w:r')
            tab = docx.oxml.OxmlElement('w:tab')
            r2.append(tab)
            new_p.append(r2)

            # Run 3: Nomor Halaman
            r3 = docx.oxml.OxmlElement('w:r')
            rPr3 = docx.oxml.OxmlElement('w:rPr')
            rf3 = docx.oxml.OxmlElement('w:rFonts')
            rf3.set(docx.oxml.ns.qn('w:ascii'), 'Times New Roman')
            rf3.set(docx.oxml.ns.qn('w:hAnsi'), 'Times New Roman')
            sz3 = docx.oxml.OxmlElement('w:sz')
            sz3.set(docx.oxml.ns.qn('w:val'), '24')
            rPr3.append(rf3)
            rPr3.append(sz3)
            t3 = docx.oxml.OxmlElement('w:t')
            t3.text = t_pg
            r3.append(rPr3)
            r3.append(t3)
            new_p.append(r3)

            parent.insert(ref_idx + offset, new_p)
            print(f"      Inserted: {t_title} -> Hlm {t_pg}")

    # =========================================================================
    # 3. Sinkronisasi DAFTAR GAMBAR (Gambar 6.1)
    # =========================================================================
    idx_dg = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "DAFTAR GAMBAR":
            idx_dg = i
            break

    if idx_dg is not None and idx_dg + 1 < len(doc.paragraphs):
        p_dg = doc.paragraphs[idx_dg + 1]
        t_nodes = list(p_dg._p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        if len(t_nodes) >= 2:
            t_nodes[0].text = "Gambar 6.1 Bagan Alur Kerangka Teoretis Penelitian"
            t_nodes[1].text = "21"
            print(f"\n   Updated DAFTAR GAMBAR: Gambar 6.1 -> Hlm 21")

    doc.save(DOCX_PATH)
    print(f"\nSUCCESS: Synced all lists to {DOCX_PATH}")

if __name__ == "__main__":
    sync_lists()

import os, sys
sys.path.append('.')
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from scripts.generate_all_instruments import (
    create_base_doc, add_header_kop, add_doc_title, add_section_heading,
    add_body_p, format_cell, set_table_borders, OUT_DIR
)

def build_instrument_04():
    doc = create_base_doc()
    add_header_kop(doc)
    add_doc_title(
        doc,
        "KISI-KISI DAN NASKAH SOAL ASESMEN DIAGNOSTIK KESIAPAN BELAJAR MATEMATIKA",
        "Materi Pokok: Rasio dan Perbandingan - Fase D (SMP Kelas VII)"
    )
    
    # Bagian A: Petunjuk Pengerjaan
    add_section_heading(doc, "A. Petunjuk Pengerjaan Tes untuk Peserta Didik")
    add_body_p(doc, "Tuliskan identitas diri Anda (Nama Lengkap, Kelas, Nomor Absen, dan Sekolah) secara jelas pada lembar jawaban yang disediakan.", "1. Identitas: ")
    add_body_p(doc, "Tes diagnostik ini bertujuan untuk memetakan pemahaman awal dan kesiapan belajar Anda pada materi Rasio, bukan untuk menentukan nilai rapor akhir semester.", "2. Tujuan: ")
    add_body_p(doc, "Kerjakan soal secara mandiri dan jujur sesuai kemampuan terbaik Anda. Jangan ragu menuliskan alur penalaran, cara coret-hitungan, atau sketsa jawaban.", "3. Kejujuran: ")
    add_body_p(doc, "Waktu pengerjaan tes adalah 60 menit. Periksa kembali seluruh jawaban sebelum diserahkan kepada guru matematika.", "4. Alokasi Waktu: ")
    
    # Bagian B: Kisi-Kisi Butir Tes
    add_section_heading(doc, "B. Kisi-Kisi Instrumen Asesmen Diagnostik")
    tbl_kisi = doc.add_table(rows=6, cols=6)
    tbl_kisi.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_kisi, "94A3B8")
    
    col_w = [Inches(0.4), Inches(0.8), Inches(3.2), Inches(0.9), Inches(0.8), Inches(0.6)]
    headers = ["No.", "Kode IK", "Indikator Ketercapaian Kompetensi", "Level Kognitif", "Bentuk Soal", "No. Soal"]
    for c_idx, h in enumerate(headers):
        tbl_kisi.rows[0].cells[c_idx].width = col_w[c_idx]
        format_cell(tbl_kisi.rows[0].cells[c_idx], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color="F1F5F9")
        
    kisi_rows = [
        ("1", "IK-01", "Menjelaskan konsep rasio dan menyatakan perbandingan dua besaran sejenis maupun berlainan satuan", "C2 (Pemahaman)", "Uraian Terstruktur", "1, 2"),
        ("2", "IK-02", "Menyederhanakan rasio dan menentukan rasio ekuivalen (senilai) pada masalah matematika", "C3 (Aplikasi)", "Uraian Terstruktur", "3, 4"),
        ("3", "IK-03", "Menentukan rasio satuan serta membandingkan laju atau harga satuan dalam kehidupan sehari-hari", "C3 (Aplikasi)", "Uraian Terstruktur", "5, 6"),
        ("4", "IK-04", "Menyelesaikan masalah kontekstual yang melibatkan perbandingan senilai dan berbalik nilai", "C3 / C4 (Penalaran)", "Uraian Kontekstual", "7, 8"),
        ("5", "IK-05", "Menyelesaikan masalah kontekstual kompleks yang melibatkan skala peta dan representasi spasial (HOTS)", "C4 (Penalaran HOTS)", "Uraian Pemecahan Masalah", "9, 10")
    ]
    for r_idx, r_vals in enumerate(kisi_rows, start=1):
        row = tbl_kisi.rows[r_idx]
        for c_idx, val in enumerate(r_vals):
            row.cells[c_idx].width = col_w[c_idx]
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 3, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(row.cells[c_idx], val, size=10, align=align)
            
    # Bagian C: Naskah Soal Tes Diagnostik
    add_section_heading(doc, "C. Naskah Butir Soal Asesmen Diagnostik")
    
    soal_items = [
        ("1", "IK-01 (Konsep Rasio)", 
         "Di kelas VII-A SMP Negeri 3 Tasikmalaya, terdapat 15 orang siswa laki-laki dan 20 orang siswa perempuan.\n"
         "a. Nyatakanlah perbandingan antara jumlah siswa laki-laki terhadap jumlah siswa perempuan dalam bentuk perbandingan paling sederhana!\n"
         "b. Nyatakanlah perbandingan antara jumlah siswa laki-laki terhadap jumlah seluruh siswa di kelas tersebut dalam bentuk pecahan biasa!"),
        
        ("2", "IK-01 (Penyetaraan Satuan)", 
         "Pak Budi memiliki dua utas kawat untuk membuat kerangka kandang kelinci. Kawat pertama memiliki panjang 2,4 meter, sedangkan kawat kedua memiliki panjang 80 sentimeter.\n"
         "Tuliskan perbandingan panjang kawat pertama terhadap panjang kawat kedua dalam bentuk perbandingan paling sederhana! Jelaskan langkah penyetaraan satuan yang Anda lakukan!"),
         
        ("3", "IK-02 (Penyederhanaan Rasio)", 
         "Ibu Nina membuat adonan sirup minuman dengan mencampurkan 350 mL konsentrat sari buah dan 1.400 mL air mineral.\n"
         "Tentukan bentuk paling sederhana dari rasio konsentrat sari buah terhadap air mineral tersebut! Tunjukkan cara Anda mencari faktor pembagi persekutuan terbesar (FPB)-nya!"),
         
        ("4", "IK-02 (Rasio Ekuivalen)", 
         "Rasio banyak kelereng merah terhadap kelereng biru di dalam sebuah toples adalah 4 : 7. Jika diketahui banyak kelereng merah adalah 28 butir:\n"
         "a. Berapakah banyak kelereng biru di dalam toples tersebut?\n"
         "b. Tentukan jumlah seluruh kelereng (merah dan biru) yang ada di dalam toples!"),
         
        ("5", "IK-03 (Laju Satuan Kecepatan)", 
         "Dua buah kendaraan melaju di jalan tol dari kota Tasikmalaya menuju Bandung:\n"
         "- Mobil A menempuh jarak sejauh 150 kilometer dalam waktu 2,5 jam.\n"
         "- Mobil B menempuh jarak sejauh 210 kilometer dalam waktu 3,5 jam.\n"
         "Hitunglah kecepatan rata-rata masing-masing mobil dalam satuan km/jam (rasio satuan), lalu tentukan mobil manakah yang melaju lebih cepat!"),
         
        ("6", "IK-03 (Harga Satuan Ekonomis)", 
         "Di sebuah minimarket, tersedia dua pilihan kemasan beras dengan merek yang sama:\n"
         "- Kemasan A: berat 2,5 kg seharga Rp35.000,00.\n"
         "- Kemasan B: berat 5 kg seharga Rp67.500,00.\n"
         "Hitunglah harga beras per kilogram pada masing-masing kemasan! Berdasarkan perhitungan rasio satuan tersebut, kemasan manakah yang lebih hemat untuk dibeli? Berikan alasan Anda!"),
         
        ("7", "IK-04 (Perbandingan Senilai)", 
         "Sebuah mobil memerlukan 6 liter bensin untuk menempuh jarak perjalanan sejauh 72 km.\n"
         "a. Berapa liter bensin yang dibutuhkan jika mobil tersebut menempuh perjalanan sejauh 180 km dengan kondisi laju yang sama?\n"
         "b. Jika tangki mobil tersebut saat ini terisi penuh sebanyak 25 liter bensin, berapa jarak maksimum yang dapat ditempuh mobil tersebut?"),
         
        ("8", "IK-04 (Perbandingan Berbalik Nilai)", 
         "Pembangunan sebuah jembatan penyeberangan direncanakan selesai dikerjakan dalam waktu 40 hari oleh 15 orang pekerja ahli.\n"
         "Karena jembatan tersebut harus segera digunakan, target penyelesaian dipercepat menjadi hanya 24 hari.\n"
         "Berapakah jumlah pekerja yang dibutuhkan seluruhnya? Berapa banyak pekerja tambahan yang harus direkrut oleh pemborong proyek tersebut?"),
         
        ("9", "IK-05 (Masalah Skala Peta - HOTS)", 
         "Pada sebuah peta provinsi Jawa Barat dengan skala 1 : 400.000, jarak lurus antara Kota Tasikmalaya dan Kota Ciamis terukur sepanjang 7,5 cm.\n"
         "a. Hitunglah jarak sebenarnya antara kedua kota tersebut dalam satuan kilometer!\n"
         "b. Jika seorang kurir ekspedisi berangkat mengendarai sepeda motor dengan kecepatan rata-rata 50 km/jam tanpa berhenti, berapa menit waktu tempuh yang dibutuhkan kurir untuk sampai di tujuan?"),
         
        ("10", "IK-05 (Masalah Denah Ruangan - HOTS)", 
         "Pak Guru merancang denah ruang laboratorium komputer sekolah berbentuk persegi panjang dengan ukuran panjang 10 cm dan lebar 6 cm pada kertas gambar berskala 1 : 120.\n"
         "a. Hitunglah ukuran panjang dan lebar sebenarnya dari ruang laboratorium tersebut dalam satuan meter!\n"
         "b. Hitunglah luas sebenarnya dari laboratorium tersebut!\n"
         "c. Jika standar kenyamanan menetapkan bahwa setiap unit meja komputer membutuhkan area seluas 3 meter persegi, berapakah kapasitas maksimal komputer yang dapat ditampung di ruangan tersebut?")
    ]
    
    for num, ik_tag, q_txt in soal_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r_num = p.add_run(f"Soal No. {num} ")
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(12)
        r_num.font.bold = True
        
        r_tag = p.add_run(f"[{ik_tag}]\n")
        r_tag.font.name = "Times New Roman"
        r_tag.font.size = Pt(11)
        r_tag.font.italic = True
        r_tag.font.color.rgb = RGBColor(30, 64, 175)
        
        r_txt = p.add_run(q_txt)
        r_txt.font.name = "Times New Roman"
        r_txt.font.size = Pt(12)
        
        # Ruang coretan jawaban
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.line_spacing = 1.15
        p_ans.paragraph_format.space_after = Pt(4)
        r_ans = p_ans.add_run("   Ruang Jawaban & Langkah Pengerjaan:\n")
        r_ans.font.name = "Times New Roman"
        r_ans.font.size = Pt(10.5)
        r_ans.font.italic = True
        r_dot = p_ans.add_run("   ............................................................................................................................................................\n   ............................................................................................................................................................")
        r_dot.font.name = "Times New Roman"
        r_dot.font.size = Pt(10.5)
        r_dot.font.color.rgb = RGBColor(148, 163, 184)

    # Bagian D: Kunci Jawaban & Rubrik Penskoran
    add_section_heading(doc, "D. Kunci Jawaban dan Rubrik Penskoran Analitis")
    add_body_p(doc, "Setiap butir soal memiliki bobot skor maksimal 10 poin, sehingga total skor sempurna adalah 100 poin. Penskoran dilakukan secara analitis berdasarkan kriteria berikut:")
    add_body_p(doc, "Jawaban kosong, tidak ada respon yang relevan, atau coretan tidak bermakna matematis.", "- Skor 0  : ")
    add_body_p(doc, "Menuliskan sebagian rumus atau konsep tetapi terdapat kesalahan fatal dalam pemahaman dasar.", "- Skor 3  : ")
    add_body_p(doc, "Alur penalaran dan konsep rumus sudah tepat, namun terjadi kekeliruan kalkulasi hitung pada hasil akhir.", "- Skor 7  : ")
    add_body_p(doc, "Langkah penyelesaian terstruktur, konseptual tepat, kalkulasi benar, dan penulisan satuan lengkap.", "- Skor 10 : ")
    
    tbl_rubrik = doc.add_table(rows=11, cols=4)
    tbl_rubrik.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_rubrik, "94A3B8")
    col_rub = [Inches(0.6), Inches(1.0), Inches(4.5), Inches(0.8)]
    headers_rub = ["No.", "Kode IK", "Kunci Jawaban & Langkah Esensial", "Skor Maks"]
    for c_idx, h in enumerate(headers_rub):
        tbl_rubrik.rows[0].cells[c_idx].width = col_rub[c_idx]
        format_cell(tbl_rubrik.rows[0].cells[c_idx], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color="F1F5F9")
        
    kunci_data = [
        ("1", "IK-01", "a. 15 : 20 = (15÷5) : (20÷5) = 3 : 4.\nb. Laki-laki / Total = 15 / (15+20) = 15/35 = 3/7.", "10"),
        ("2", "IK-01", "Penyetaraan: 2,4 m = 2,4 × 100 cm = 240 cm.\nRasio: 240 cm : 80 cm = (240÷80) : (80÷80) = 3 : 1.", "10"),
        ("3", "IK-02", "Rasio: 350 mL : 1.400 mL. FPB(350, 1.400) = 350.\nBentuk sederhana: (350÷350) : (1.400÷350) = 1 : 4.", "10"),
        ("4", "IK-02", "a. 4/7 = 28/B => B = (7 × 28) ÷ 4 = 196 ÷ 4 = 49 butir.\nb. Total kelereng = 28 + 49 = 77 butir.", "10"),
        ("5", "IK-03", "Mobil A: 150 km ÷ 2,5 jam = 60 km/jam.\nMobil B: 210 km ÷ 3,5 jam = 60 km/jam.\nKesimpulan: Kedua mobil melaju dengan kecepatan rata-rata sama (60 km/jam).", "10"),
        ("6", "IK-03", "Kemasan A: Rp35.000 ÷ 2,5 kg = Rp14.000 / kg.\nKemasan B: Rp67.500 ÷ 5 kg = Rp13.500 / kg.\nKesimpulan: Kemasan B lebih hemat (selisih lebih murah Rp500 per kg).", "10"),
        ("7", "IK-04", "a. Konsumsi per km = 72÷6 = 12 km/L. Bensin untuk 180 km = 180÷12 = 15 liter.\nb. Jarak tempuh 25 liter = 25 × 12 km = 300 km.", "10"),
        ("8", "IK-04", "Perbandingan berbalik nilai: 40 × 15 = 24 × x => 600 = 24x => x = 25 pekerja.\nPekerja tambahan = 25 - 15 = 10 orang pekerja tambahan.", "10"),
        ("9", "IK-05", "a. Jarak sebenarnya = 7,5 cm × 400.000 = 3.000.000 cm = 30 km.\nb. Waktu = Jarak ÷ Kecepatan = 30 km ÷ 50 km/jam = 0,6 jam = 0,6 × 60 menit = 36 menit.", "10"),
        ("10", "IK-05", "a. P = 10 × 120 cm = 1.200 cm = 12 m. L = 6 × 120 cm = 720 cm = 7,2 m.\nb. Luas sebenarnya = 12 m × 7,2 m = 86,4 m².\nc. Kapasitas meja komputer = 86,4 ÷ 3 = 28,8 => maksimum 28 unit komputer.", "10")
    ]
    
    for r_idx, r_vals in enumerate(kunci_data, start=1):
        row = tbl_rubrik.rows[r_idx]
        for c_idx, val in enumerate(r_vals):
            row.cells[c_idx].width = col_rub[c_idx]
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 3] else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(row.cells[c_idx], val, size=9.5, align=align)

    # Bagian E: Aturan Klasifikasi Kesiapan Belajar
    add_section_heading(doc, "E. Aturan Algoritma Klasifikasi Kesiapan Belajar Siswa")
    add_body_p(doc, "Pengelompokan kesiapan belajar peserta didik ke dalam tiga kategori kurikuler resmi (Pedoman Pembelajaran dan Asesmen BSKAP 2024) dilakukan dengan aturan berbasis bukti (evidence-centered design) berikut:")
    
    tbl_klas = doc.add_table(rows=4, cols=4)
    tbl_klas.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_klas, "94A3B8")
    col_klas = [Inches(1.8), Inches(1.3), Inches(2.7), Inches(1.3)]
    headers_klas = ["Kategori Kesiapan", "Rentang Skor", "Karakteristik Penguasaan Indikator", "Rekomendasi Varian LKPD"]
    for c_idx, h in enumerate(headers_klas):
        tbl_klas.rows[0].cells[c_idx].width = col_klas[c_idx]
        format_cell(tbl_klas.rows[0].cells[c_idx], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color="F1F5F9")
        
    klas_rows = [
        ("Perlu Bimbingan\n(Tier 1 - Dasar)", "0 – 59 Poin\n(< 60%)", "Belum menguasai konsep dasar rasio (IK-01/IK-02) atau mengalami miskonsepsi aditif keliru.", "LKPD Varian A\n(Scaffolding Tinggi & Terstruktur)"),
        ("Berkembang\n(Tier 2 - Menengah)", "60 – 79 Poin\n(60% s.d. 79%)", "Menguasai konsep dasar (IK-01, IK-02, IK-03), namun masih memerlukan bantuan pada soal kontekstual kompleks.", "LKPD Varian B\n(Fading Guidance & Pendampingan Sedang)"),
        ("Mahir\n(Tier 3 - Lanjut)", "80 – 100 Poin\n(≥ 80%)", "Menguasai seluruh indikator kompetensi (IK-01 s.d. IK-05) serta mampu menyelesaikan soal HOTS secara mandiri.", "LKPD Varian C\n(Tantangan HOTS & Eksplorasi Mandiri)")
    ]
    for r_idx, r_vals in enumerate(klas_rows, start=1):
        row = tbl_klas.rows[r_idx]
        for c_idx, val in enumerate(r_vals):
            row.cells[c_idx].width = col_klas[c_idx]
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 3] else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(row.cells[c_idx], val, size=10, align=align)
            
    out_docx = os.path.join(OUT_DIR, "04_Kisi_Kisi_dan_Naskah_Soal_Tes_Diagnostik.docx")
    doc.save(out_docx)
    print("Saved:", out_docx)

if __name__ == "__main__":
    build_instrument_04()

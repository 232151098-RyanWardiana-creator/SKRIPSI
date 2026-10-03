import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_base_document():
    doc = docx.Document()
    # Page setup: A4, Margins: Left 3cm, Right 2.5cm, Top 2.5cm, Bottom 2.5cm
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1.18)   # 3.0 cm
        section.right_margin = Inches(0.98)  # 2.5 cm
        section.top_margin = Inches(0.98)    # 2.5 cm
        section.bottom_margin = Inches(0.98) # 2.5 cm
        section.different_first_page_header_footer = False
    
    # Set default style to Times New Roman 12pt
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.space_before = Pt(0)
    
    return doc

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)

def add_header_kop(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(2)
    
    r1 = p.add_run("UNIVERSITAS SILIWANGI\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(14)
    r1.font.bold = True
    
    r2 = p.add_run("FAKULTAS KEGURUAN DAN ILMU PENDIDIKAN\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    r2.font.bold = True
    
    r3 = p.add_run("JURUSAN PENDIDIKAN MATEMATIKA\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True
    
    r4 = p.add_run("Jalan Siliwangi No. 24 Kota Tasikmalaya, Jawa Barat 46115\nLaman: https://fkip.unsil.ac.id | Pos-el: mat@unsil.ac.id")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(9.5)
    r4.font.italic = True
    
    # Border line below kop
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(2)
    p_line.paragraph_format.space_after = Pt(12)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="double" w:sz="12" w:space="1" w:color="000000"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(pBdr)

def add_doc_title(doc, title_text, subtitle_text=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    
    r = p.add_run(title_text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    
    if subtitle_text:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.line_spacing = 1.15
        p2.paragraph_format.space_after = Pt(14)
        r2 = p2.add_run(subtitle_text)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)
        r2.font.italic = True

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

def add_body_p(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.italic = italic
    return p

print("Core helpers ready!")

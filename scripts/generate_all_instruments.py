import os, sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client
import pythoncom

OUT_DIR = r"C:\Users\LATITUDE 3310 TOUCH\Documents\DOCUMENT RYAN\02_Kuliah\Semester_7\KF21518001 - Skripsi (Kelas A)\04_Instrumen dan Validasi Soal-Angket"
os.makedirs(OUT_DIR, exist_ok=True)

def create_base_doc():
    doc = docx.Document()
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.left_margin = Inches(1.18)   # 3.0 cm
        s.right_margin = Inches(0.98)  # 2.5 cm
        s.top_margin = Inches(0.98)    # 2.5 cm
        s.bottom_margin = Inches(0.98) # 2.5 cm
    
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.space_before = Pt(0)
    return doc

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_table_borders(table, color="B0BEC5", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)

def add_header_kop(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(2)
    
    r1 = p.add_run("UNIVERSITAS SILIWANGI\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(14)
    r1.font.bold = True
    
    r2 = p.add_run("FAKULTAS KEGURUAN DAN ILMU PENDIDIKAN\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    r2.font.bold = True
    
    r3 = p.add_run("JURUSAN PENDIDIKAN MATEMATIKA\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True
    
    r4 = p.add_run("Jalan Siliwangi No. 24 Kota Tasikmalaya, Jawa Barat 46115\nLaman: https://fkip.unsil.ac.id | Pos-el: mat@unsil.ac.id")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(9.5)
    r4.font.italic = True
    
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(2)
    p_line.paragraph_format.space_after = Pt(12)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="double" w:sz="12" w:space="1" w:color="000000"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(pBdr)

def add_doc_title(doc, title_text, subtitle_text=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    
    r = p.add_run(title_text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    
    if subtitle_text:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.line_spacing = 1.15
        p2.paragraph_format.space_after = Pt(14)
        r2 = p2.add_run(subtitle_text)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)
        r2.font.italic = True

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

def add_body_p(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.italic = italic
    return p

def format_cell(cell, text="", bold=False, italic=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, v_align=WD_ALIGN_VERTICAL.CENTER, color=None):
    set_cell_margins(cell)
    cell.vertical_alignment = v_align
    if color:
        set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic

print("Base setup verified!")

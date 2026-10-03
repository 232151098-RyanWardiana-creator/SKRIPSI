import os
import shutil
import zipfile
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

MASTER_DOCX = r"C:\Users\LATITUDE 3310 TOUCH\Documents\DOCUMENT RYAN\02_Kuliah\Semester_7\KF21518001 - Skripsi (Kelas A)\03_Draft Proposal (BAB 1-3)\Proposal_Skripsi_Ryan_Wardiana_Revisi_Bimbingan_2_Draft.docx"
BACKUP_DOCX = r"C:\Users\LATITUDE 3310 TOUCH\Documents\DOCUMENT RYAN\02_Kuliah\Semester_7\KF21518001 - Skripsi (Kelas A)\03_Draft Proposal (BAB 1-3)\Proposal_Skripsi_Ryan_Wardiana_Revisi_Bimbingan_2_Draft.docx.bak"
IMAGE_SRC = r"C:\Users\LATITUDE 3310 TOUCH\Documents\DOCUMENT RYAN\02_Kuliah\Semester_7\KF21518001 - Skripsi (Kelas A)\03_Draft Proposal (BAB 1-3)\gambar_6_1_kerangka_teoretis.png"

def italicize_terms_in_paragraph(p, terms):
    """Italicizes specific English/foreign terms in a paragraph while preserving font."""
    full_text = p.text
    if not any(term in full_text for term in terms):
        return
    sorted_terms = sorted(terms, key=len, reverse=True)
    pattern = '(' + '|'.join(re.escape(term) for term in sorted_terms) + ')'
    parts = re.split(pattern, full_text)
    
    alignment = p.alignment
    left_indent = p.paragraph_format.left_indent
    first_line_indent = p.paragraph_format.first_line_indent
    space_before = p.paragraph_format.space_before
    space_after = p.paragraph_format.space_after
    line_spacing = p.paragraph_format.line_spacing
    
    p.text = ""
    for part in parts:
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0, 0, 0)
        if part in terms:
            r.italic = True
            
    p.alignment = alignment
    p.paragraph_format.left_indent = left_indent
    p.paragraph_format.first_line_indent = first_line_indent
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = line_spacing

def run_reorganization():
    # 1. Pastikan backup tersedia
    if not os.path.exists(BACKUP_DOCX):
        shutil.copy2(MASTER_DOCX, BACKUP_DOCX)
        print("Backup created at:", BACKUP_DOCX)

    # 2. Ganti Gambar 6.1 pada word/media/image2.png
    temp_zip = MASTER_DOCX + ".tmp.zip"
    with zipfile.ZipFile(MASTER_DOCX, 'r') as zin:
        with zipfile.ZipFile(temp_zip, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "word/media/image2.png":
                    with open(IMAGE_SRC, 'rb') as f_img:
                        zout.writestr(item, f_img.read())
                    print("2. Gambar 6.1 berhasil diperbarui di word/media/image2.png!")
                else:
                    zout.writestr(item, zin.read(item.filename))
    os.replace(temp_zip, MASTER_DOCX)

    doc = docx.Document(MASTER_DOCX)
    body = doc._element.body
    paras = doc.paragraphs

    # =========================================================
    # PART A: REORDER DEFINISI OPERASIONAL (BAB I)
    # =========================================================
    p71_node = next(p for p in paras if p.text.strip() == "Tujuan Penelitian")._p

    def get_do_pair_by_title(title_text):
        idx = next(i for i, p in enumerate(paras) if p.style.name == "Heading 2" and p.text.strip() == title_text)
        return [paras[idx]._p, paras[idx + 1]._p]

    pair_rasio = get_do_pair_by_title("Materi Rasio dan Perbandingan")
    pair_diagnostik = get_do_pair_by_title("Asesmen Diagnostik dan Kesiapan Belajar")
    pair_tarl = get_do_pair_by_title("Pembelajaran Berdiferensiasi dan Teaching at the Right Level (TaRL)")
    pair_scaffolding = get_do_pair_by_title("Scaffolding dan Higher Order Thinking Skills (HOTS)")
    pair_genai = get_do_pair_by_title("Generative Artificial Intelligence (GenAI) dan Large Language Model (LLM)")
    pair_webapp = get_do_pair_by_title("Web Application LKPD Berbasis GenAI")
    pair_omml = get_do_pair_by_title("Representasi Formula Matematika dan Transformasi ke Office Math Markup Language (OMML)")

    all_do_pairs = [
        pair_rasio, pair_diagnostik, pair_tarl,
        pair_scaffolding, pair_genai, pair_webapp, pair_omml
    ]

    for pair in all_do_pairs:
        for node in pair:
            body.remove(node)

    target_do_pairs = [
        pair_webapp,
        pair_genai,
        pair_omml,
        pair_tarl,
        pair_diagnostik,
        pair_scaffolding,
        pair_rasio
    ]

    insert_idx = body.index(p71_node)
    for pair in target_do_pairs:
        for node in pair:
            body.insert(insert_idx, node)
            insert_idx += 1
    print("3. Definisi Operasional berhasil diurutkan ulang sesuai alur judul!")

    # =========================================================
    # PART B: REORDER KAJIAN TEORI (BAB II - Subbab 6.1)
    # =========================================================
    def get_kt_block_by_title(title_text):
        idx = next(i for i, p in enumerate(paras) if p.style.name == "Heading 3" and p.text.strip() == title_text)
        return [paras[idx + k]._p for k in range(5)]

    block_rasio = get_kt_block_by_title("Materi Rasio dan Perbandingan")
    block_diagnostik = get_kt_block_by_title("Asesmen Diagnostik dan Kesiapan Belajar")
    block_tarl = get_kt_block_by_title("Pembelajaran Berdiferensiasi dan Teaching at the Right Level (TaRL)")
    block_scaffolding = get_kt_block_by_title("Scaffolding dan Higher Order Thinking Skills (HOTS)")
    block_genai = get_kt_block_by_title("Generative Artificial Intelligence (GenAI) dan Large Language Model (LLM)")
    block_webapp = get_kt_block_by_title("LKPD Elektronik Berbasis Web dan Model Pengembangan ADDIE")
    block_omml = get_kt_block_by_title("Representasi Formula Matematika dan Transformasi ke Office Math Markup Language (OMML)")

    all_kt_blocks = [
        block_rasio, block_diagnostik, block_tarl,
        block_scaffolding, block_genai, block_webapp, block_omml
    ]

    idx_conc = next(i for i, p in enumerate(paras) if "Sintesis di atas menunjukkan empat celah ilmiah" in p.text)
    p_spacer_kt = paras[idx_conc - 1]

    for blk in all_kt_blocks:
        for node in blk:
            body.remove(node)

    target_kt_blocks = [
        block_webapp,
        block_genai,
        block_omml,
        block_tarl,
        block_diagnostik,
        block_scaffolding,
        block_rasio
    ]

    insert_kt_idx = body.index(p_spacer_kt._p)
    for blk in target_kt_blocks:
        for node in blk:
            body.insert(insert_kt_idx, node)
            insert_kt_idx += 1
    print("4. Kajian Teori (BAB II) berhasil disinkronkan persis urutan DO!")

    # =========================================================
    # PART C: RESTRUKTURISASI SUBBAB 7.5 (TEKNIK ANALISIS DATA) & REPOSISI TABEL
    # =========================================================
    p167 = next(p for p in paras if p.text.strip() == "Teknik Analisis Data")
    idx_168 = next(i for i, p in enumerate(paras) if "1. Analisis Validitas Isi Instrumen Asesmen Diagnostik" in p.text)
    p168 = paras[idx_168]
    p169 = paras[idx_168 + 1] # Narasi Aiken

    p173 = next(p for p in paras if "2. Analisis Kelayakan Produk LKPD" in p.text)
    p177 = next(p for p in paras if "Kriteria kevalidan produk mengacu pada kriteria kelayakan" in p.text)

    p178 = next(p for p in paras if "3. Analisis Kepraktisan Penggunaan Produk" in p.text)
    p182 = next(p for p in paras if "Produk dinyatakan praktis dan siap diimplementasikan" in p.text)

    # Identifikasi caption dan tabel asli di Subbab 7.5 HANYA yang ber-style 'Caption'
    p_cap710 = next(p for p in paras if p.style.name == "Caption" and "Tabel 7.10" in p.text)
    idx_cap710 = body.index(p_cap710._p)
    tbl_710 = body[idx_cap710 + 1]
    spacer_710 = body[idx_cap710 + 2]

    p_cap711 = next(p for p in paras if p.style.name == "Caption" and "Tabel 7.11" in p.text)
    idx_cap711 = body.index(p_cap711._p)
    tbl_711 = body[idx_cap711 + 1]
    spacer_711 = body[idx_cap711 + 2]

    # Ubah p168 menjadi judul besar Analisis Kevalidan
    p168.text = "1. Analisis Kevalidan (Validitas Produk dan Instrumen)"
    p168.runs[0].bold = True

    # Sisipkan narasi pengantar kevalidan
    intro_kevalidan = p169.insert_paragraph_before(
        "Analisis kevalidan dilakukan untuk memastikan kelayakan teoritis dan empiris instrumen asesmen serta produk perangkat ajar yang dikembangkan. Analisis ini mencakup dua dimensi evaluasi terpadu, yaitu validitas isi butir instrumen asesmen diagnostik oleh panel ahli serta validitas kelayakan produk web application dan LKPD oleh ahli materi dan ahli media."
    )
    intro_kevalidan.style = "Paragraph"
    intro_kevalidan.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_kevalidan.paragraph_format.first_line_indent = Inches(0.4)
    intro_kevalidan.paragraph_format.line_spacing = 1.5
    intro_kevalidan.paragraph_format.space_before = Pt(3)
    intro_kevalidan.paragraph_format.space_after = Pt(4)

    # Sub-poin a
    sub_a_title = p169.insert_paragraph_before("a. Validitas Isi Butir Instrumen Asesmen Diagnostik (Koefisien Aiken's V)")
    sub_a_title.style = "Paragraph"
    sub_a_title.runs[0].bold = True
    sub_a_title.paragraph_format.left_indent = Inches(0.2)
    sub_a_title.paragraph_format.space_before = Pt(4)
    sub_a_title.paragraph_format.space_after = Pt(2)

    # Sub-poin b (mengubah p173)
    p173.text = "b. Validitas Kelayakan Produk LKPD dan Web Application"
    p173.runs[0].bold = True
    p173.paragraph_format.left_indent = Inches(0.2)
    p173.paragraph_format.space_before = Pt(6)
    p173.paragraph_format.space_after = Pt(2)

    # Sesuaikan teks kriteria Akbar
    p177.text = "Kriteria kevalidan produk mengacu pada kriteria kelayakan Sa'dun Akbar (2013, hlm. 83) sebagaimana disajikan pada Tabel 7.10 berikut."

    # Ubah p178 menjadi Poin 2
    p178.text = "2. Analisis Kepraktisan Penggunaan Produk (Respon Guru dan Siswa)"
    p178.runs[0].bold = True
    p178.paragraph_format.left_indent = Inches(0)
    p178.paragraph_format.space_before = Pt(8)
    p178.paragraph_format.space_after = Pt(2)

    # Sesuaikan teks kriteria kepraktisan p182
    p182.text = "Produk dinyatakan praktis dan siap diimplementasikan apabila persentase kepraktisan mencapai minimal kategori 'Praktis' mengacu pada Tabel 7.11 berikut. Kriteria interval tersebut diadaptasi dari Riduwan (2006, hlm. 89) sebagaimana dirujuk Erlinda dan Lelfita (2020)."

    # PINDAHKAN TABEL 7.10 TEPAT DI BAWAH p177 (sebelum p178):
    body.remove(p_cap710._p)
    body.remove(tbl_710)
    body.remove(spacer_710)

    idx_target_710 = body.index(p178._p)
    body.insert(idx_target_710, p_cap710._p)
    body.insert(idx_target_710 + 1, tbl_710)
    body.insert(idx_target_710 + 2, spacer_710)

    # PINDAHKAN TABEL 7.11 TEPAT DI BAWAH p182 (sebelum Heading 2 Tempat dan Jadwal Penelitian):
    body.remove(p_cap711._p)
    body.remove(tbl_711)
    body.remove(spacer_711)

    p_waktu_h = next(p for p in paras if "Tempat dan Jadwal Penelitian" in p.text)
    idx_target_711 = body.index(p_waktu_h._p)
    body.insert(idx_target_711, p_cap711._p)
    body.insert(idx_target_711 + 1, tbl_711)
    body.insert(idx_target_711 + 2, spacer_711)
    print("5. Restrukturisasi Subbab 7.5 dan penempatan Tabel 7.10 & 7.11 berhasil!")

    # =========================================================
    # PART D: PEMISAHAN & PENAMBAHAN 5 ALASAN AKADEMIS SUBBAB 7.6
    # =========================================================
    p_waktu_h.text = "Waktu dan Tempat Penelitian" # Heading 2

    idx_wh = next(i for i, p in enumerate(paras) if p._p == p_waktu_h._p)
    p188 = paras[idx_wh + 1]
    p188.text = "" # Kosongkan teks lama

    # 1. Tempat Penelitian
    p_sub1 = p188.insert_paragraph_before()
    p_sub1.style = "Paragraph"
    r_sub1 = p_sub1.add_run("1. Tempat Penelitian")
    r_sub1.bold = True
    p_sub1.paragraph_format.space_before = Pt(6)
    p_sub1.paragraph_format.space_after = Pt(2)

    p_loc = p188.insert_paragraph_before(
        "Penelitian dan pengembangan ini dilaksanakan di dua lokasi utama, yaitu: (1) Jurusan Pendidikan Matematika, Fakultas Keguruan dan Ilmu Pendidikan, Universitas Siliwangi sebagai basis perancangan konseptual, rekayasa perangkat lunak sistem, dan pengujian teknis; serta (2) SMP Negeri 3 Tasikmalaya (beralamat di Jl. Merdeka No. 17, Tawangsari, Kec. Tawang, Kota Tasikmalaya) sebagai lokasi studi pendahuluan, validasi pengguna, dan uji coba implementasi empiris."
    )
    p_loc.style = "Paragraph"
    p_loc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_loc.paragraph_format.first_line_indent = Inches(0.4)
    p_loc.paragraph_format.line_spacing = 1.5
    p_loc.paragraph_format.space_before = Pt(2)
    p_loc.paragraph_format.space_after = Pt(3)

    p_reason_intro = p188.insert_paragraph_before(
        "Pemilihan SMP Negeri 3 Tasikmalaya sebagai lokasi penelitian didasarkan pada lima pertimbangan akademis dan metodologis yang objektif:"
    )
    p_reason_intro.style = "Paragraph"
    p_reason_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_reason_intro.paragraph_format.first_line_indent = Inches(0.4)
    p_reason_intro.paragraph_format.line_spacing = 1.5
    p_reason_intro.paragraph_format.space_before = Pt(2)
    p_reason_intro.paragraph_format.space_after = Pt(2)

    foreign_terms = [
        "web application", "Generative AI", "Office Math Markup Language", "OMML",
        "Teaching at the Right Level", "Bring Your Own Device", "smartphone",
        "scaffolding", "fading guidance", "Zone of Proximal Development", "ADDIE",
        "Analysis", "Design", "Development", "Implementation", "Evaluation",
        "Feasibility and Intensive Monitoring", "human-in-the-loop"
    ]

    reasons = [
        ("a. Implementasi Mandiri Kurikulum Merdeka Fase D: ",
         "SMP Negeri 3 Tasikmalaya telah mengimplementasikan Kurikulum Merdeka secara aktif pada jenjang kelas VII (Fase D), yang secara regulasi menuntut pelaksanaan asesmen awal pembelajaran (asesmen diagnostik kognitif) serta realisasi pembelajaran berdiferensiasi (Teaching at the Right Level) untuk melayani keberagaman kebutuhan belajar peserta didik."),
        ("b. Kebutuhan Nyata Guru Matematika Terhadap Bahan Ajar Berdiferensiasi: ",
         "Berdasarkan hasil studi pendahuluan dengan guru matematika kelas VII, terdapat kebutuhan mendesak akan inovasi bahan ajar adaptif. Guru menghadapi kendala alokasi waktu yang signifikan jika harus merancang tiga varian bahan ajar secara manual serta menghadapi hambatan teknis pengetikan formula matematika pada pengolah kata standar. Kehadiran web application LKPD-AI berbasis Generative AI dan OMML menjadi solusi relevan atas permasalahan empiris tersebut."),
        ("c. Heterogenitas Kesiapan Belajar Siswa pada Materi Rasio: ",
         "Peserta didik kelas VII SMP Negeri 3 Tasikmalaya berasal dari sebaran sekolah dasar yang beragam dengan rentang penguasaan materi prasyarat matematika (pecahan, desimal, dan faktor persekutuan) yang bervariasi secara signifikan. Keragaman kemampuan awal ini menjadi subjek empiris yang tepat dan valid untuk menguji keterbacaan serta keefektifan scaffolding tiga tingkatan (Perlu Bimbingan, Berkembang, dan Mahir)."),
        ("d. Optimalisasi Literasi Digital melalui Kebijakan Gawai Mandiri (Bring Your Own Device): ",
         "Sekolah menerapkan kebijakan yang mendukung pemanfaatan teknologi secara positif dan edukatif di lingkungan kelas, di mana peserta didik diperkenankan menggunakan gawai mandiri (smartphone) secara terarah di bawah bimbingan guru. Kebijakan ini mendukung operasionalisasi sistem web application, di mana siswa dapat mengakses instrumen asesmen diagnostik daring dan dasbor pengisian LKPD interaktif secara langsung dan fleksibel tanpa bergantung penuh pada laboratorium komputer sekolah, sekaligus melatih kemandirian dan literasi digital siswa."),
        ("e. Kelayakan Pemantauan Intensif (Feasibility and Intensive Monitoring): ",
         "Lokasi sekolah yang strategis di pusat Kota Tasikmalaya dan berada dalam radius jangkauan yang sangat dekat dengan kampus Universitas Siliwangi menjamin efisiensi koordinasi teknis antara peneliti, guru mitra, dan validator. Hal ini memungkinkan pemantauan observasi kelas, pendampingan uji coba produk, serta konfirmasi umpan balik siklus ADDIE secara intensif, berkelanjutan, dan mendalam.")
    ]

    pattern = '(' + '|'.join(re.escape(term) for term in sorted(foreign_terms, key=len, reverse=True)) + ')'
    for label, desc in reasons:
        p_r = p188.insert_paragraph_before()
        p_r.style = "Paragraph"
        p_r.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_r.paragraph_format.left_indent = Inches(0.4)
        p_r.paragraph_format.line_spacing = 1.5
        p_r.paragraph_format.space_before = Pt(2)
        p_r.paragraph_format.space_after = Pt(2)
        r_lbl = p_r.add_run(label)
        r_lbl.bold = True
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(12)
        
        parts = re.split(pattern, desc)
        for part in parts:
            if not part:
                continue
            r = p_r.add_run(part)
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            if part in foreign_terms:
                r.italic = True

    # 2. Waktu dan Jadwal Penelitian
    p_sub2 = p188.insert_paragraph_before()
    p_sub2.style = "Paragraph"
    r_sub2 = p_sub2.add_run("2. Waktu dan Jadwal Penelitian")
    r_sub2.bold = True
    p_sub2.paragraph_format.space_before = Pt(6)
    p_sub2.paragraph_format.space_after = Pt(2)

    p_time = p188.insert_paragraph_before(
        "Penelitian dan pengembangan ini direncanakan berlangsung selama enam bulan, terhitung mulai bulan September 2026 sampai dengan Februari 2027, mencakup lima fase model pengembangan ADDIE (Analysis, Design, Development, Implementation, dan Evaluation). Rincian matriks jadwal pelaksanaan penelitian disajikan pada Tabel 7.12 berikut."
    )
    p_time.style = "Paragraph"
    p_time.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_time.paragraph_format.first_line_indent = Inches(0.4)
    p_time.paragraph_format.line_spacing = 1.5
    p_time.paragraph_format.space_before = Pt(2)
    p_time.paragraph_format.space_after = Pt(4)

    # Hapus p188 kosong dari body
    body.remove(p188._p)

    new_paragraphs = [intro_kevalidan, p_loc, p_reason_intro, p_time]
    for p in new_paragraphs:
        italicize_terms_in_paragraph(p, foreign_terms)

    doc.save(MASTER_DOCX)
    print("6. Naskah Word MASTER_DOCX berhasil diperbarui!")

if __name__ == "__main__":
    run_reorganization()

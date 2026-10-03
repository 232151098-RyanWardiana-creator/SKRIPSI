import os, sys
sys.path.append('.')
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from scripts.generate_all_instruments import (
    create_base_doc, add_header_kop, add_doc_title, add_section_heading,
    add_body_p, format_cell, set_table_borders, OUT_DIR
)

def build_instrument_06():
    doc = create_base_doc()
    add_header_kop(doc)
    add_doc_title(
        doc,
        "ANGKET RESPON KEPRAKTISAN PENGGUNAAN OLEH GURU MATEMATIKA",
        "Uji Coba Lapangan Web Application Berbasis AI untuk Menghasilkan LKPD Matematika Berdiferensiasi pada Materi Rasio"
    )
    
    # Bagian A: Identitas Responden Guru
    add_section_heading(doc, "A. Identitas Responden Guru Matematika")
    tbl_id = doc.add_table(rows=6, cols=3)
    tbl_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_id, "E2E8F0")
    col_widths = [Inches(0.4), Inches(2.3), Inches(3.8)]
    fields = [
        ("1", "Nama Lengkap & Gelar", ": ....................................................................................."),
        ("2", "NIP / NUPTK", ": ....................................................................................."),
        ("3", "Sekolah / Unit Kerja", ": SMP Negeri 3 Tasikmalaya / Sekolah Mitra"),
        ("4", "Kelas yang Diampu", ": Kelas VII (Fase D)"),
        ("5", "Masa Kerja / Pengalaman Mengajar", ": ............. Tahun"),
        ("6", "Hari, Tanggal Pengisian", ": .....................................................................................")
    ]
    for r_idx, (no, label, val) in enumerate(fields):
        row = tbl_id.rows[r_idx]
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = w
        format_cell(row.cells[0], no, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[1], label, bold=True, size=11)
        format_cell(row.cells[2], val, size=11)
        
    # Bagian B: Petunjuk Pengisian
    add_section_heading(doc, "B. Petunjuk Pengisian Angket Kepraktisan")
    add_body_p(doc, "Angket ini bertujuan untuk mengukur tingkat kepraktisan, kemudahan operasional, efisiensi waktu, dan kebermanfaatan nyata web application berbasis AI yang dikembangkan dalam membantu Bapak/Ibu merancang LKPD berdiferensiasi.", "1. Tujuan: ")
    add_body_p(doc, "Bapak/Ibu dimohon memberikan tanda centang (✓) pada salah satu pilihan jawaban yang paling sesuai dengan pengalaman langsung Bapak/Ibu setelah mencoba menggunakan sistem.", "2. Cara Menjawab: ")
    add_body_p(doc, "Keterangan pilihan jawaban: 5 = Sangat Setuju (SS); 4 = Setuju (S); 3 = Cukup Setuju / Netral (CS); 2 = Tidak Setuju (TS); 1 = Sangat Tidak Setuju (STS).", "3. Kategori Skala: ")
    add_body_p(doc, "Seluruh jawaban dan masukan yang Bapak/Ibu berikan semata-mata dimanfaatkan untuk kepentingan pengembangan akademik skripsi.", "4. Kerahasiaan: ")

    # Bagian C: Tabel 30 Butir (Tabel 7.5)
    add_section_heading(doc, "C. Lembar Penilaian Kepraktisan Penggunaan")
    
    items_data = [
        # Aspek 1: Kemudahan Dipelajari (1-4)
        ("I. Aspek Kemudahan Dipelajari (Learnability)", [
            ("1", "Tujuan, fungsi utama, dan alur operasional web application ini sangat mudah saya pahami saat pertama kali dijelaskan."),
            ("2", "Tata letak tombol, navigasi menu, dan instruksi pada antarmuka aplikasi tidak membingungkan pengguna pemula."),
            ("3", "Istilah matematika, ikon grafis, dan label fitur yang digunakan terasa familier dan komunikatif bagi guru."),
            ("4", "Panduan penggunaan ringkas yang disediakan di dalam aplikasi sangat membantu mempercepat pemahaman operasional sistem.")
        ]),
        # Aspek 2: Kemudahan Digunakan (5-10)
        ("II. Aspek Kemudahan Digunakan (Operability)", [
            ("5", "Proses mendistribusikan tautan asesmen diagnostik dan mengumpulkan jawaban peserta didik sangat praktis dan lancar."),
            ("6", "Halaman dasbor ringkasan profil kesiapan belajar peserta didik per indikator mudah dibaca dan diinterpretasikan."),
            ("7", "Proses memilih topik materi rasio, menentukan capaian, dan memilih jenis diferensiasi berjalan mudah."),
            ("8", "Proses memicu generasi tiga draf varian LKPD secara bersamaan dengan tombol aksi AI sangat sederhana dan cepat."),
            ("9", "Fasilitas editor layar untuk meninjau, menyisipkan teks tambahan, atau mengoreksi butir LKPD sangat mudah dioperasikan."),
            ("10", "Perpindahan antarfitur pada aplikasi terasa responsif, intuitif, dan tidak mengalami hambatan teknis.")
        ]),
        # Aspek 3: Efisiensi Waktu (11-13)
        ("III. Aspek Efisiensi Alokasi Waktu Kerja Guru", [
            ("11", "Penggunaan web application ini secara nyata memangkas durasi waktu persiapan merancang perangkat pembelajaran di sekolah."),
            ("12", "Proses pengolahan data diagnostik dan pengelompokan siswa otomatis jauh lebih cepat dibandingkan analisis manual."),
            ("13", "Ketersediaan draf tiga varian LKPD (Perlu Bimbingan, Berkembang, Mahir) sekaligus sangat meringankan beban administratif guru.")
        ]),
        # Aspek 4: Kejelasan dan Kualitas Keluaran (14-18)
        ("IV. Aspek Kejelasan dan Kualitas Keluaran LKPD", [
            ("14", "Pengelompokan kesiapan siswa ke dalam 3 tingkat (Perlu Bimbingan, Berkembang, Mahir) sangat akurat dan transparan."),
            ("15", "Susunan urutan kegiatan pada masing-masing draf varian LKPD runtut, logis, dan sesuai dengan tahapan belajar peserta didik."),
            ("16", "Perbedaan gradasi bantuan belajar (scaffolding) pada ketiga varian LKPD terlihat kontras, proporsional, dan tepat sasaran."),
            ("17", "Konteks masalah nyata yang disajikan dalam lembar kerja menarik, realistis, dan relevan dengan lingkungan peserta didik."),
            ("18", "Tampilan narasi soal, tabel data, ilustrasi pendukung, dan rumus matematika tersaji dengan jelas dan mudah dipahami.")
        ]),
        # Aspek 5: Kontrol dan Fleksibilitas Pengguna (19-22)
        ("V. Aspek Kontrol dan Fleksibilitas Guru (Human-in-the-Loop)", [
            ("19", "Aplikasi memberikan ruang kendali penuh bagi guru untuk menelaah dan menyunting teks draf keluaran kecerdasan artifisial."),
            ("20", "Guru memiliki keleluasaan untuk membatalkan proses, mengulang pembuatan draf, atau mengganti stimulus soal."),
            ("21", "Guru tetap bertindak sebagai pemegang keputusan pedagogis tertinggi sebelum draf lembar kerja dibagikan ke siswa."),
            ("22", "Sistem tidak memaksakan ekspor otomatis tanpa adanya verifikasi dan persetujuan (approval) langsung dari guru.")
        ]),
        # Aspek 6: Kemudahan Ekspor dan Pencetakan (23-26)
        ("VI. Aspek Kemudahan Ekspor dan Pencetakan Dokumen Word", [
            ("23", "Proses mengunduh draf LKPD menjadi berkas dokumen Microsoft Word (.docx) berlangsung lancar dan tanpa kendala."),
            ("24", "Berkas Word hasil unduhan langsung dapat dibuka secara normal di komputer tanpa menimbulkan pesan peringatan korup."),
            ("25", "Formula matematika di dalam berkas Word hasil ekspor dapat diedit secara langsung menggunakan Equation Editor bawaan (OMML)."),
            ("26", "Tata letak, batas margin, dan kerapian halaman berkas Word langsung siap cetak (print-ready) pada kertas A4.")
        ]),
        # Aspek 7: Kepuasan dan Kebermanfaatan (27-30)
        ("VII. Aspek Kepuasan dan Kebermanfaatan Pembelajaran", [
            ("27", "Merasa sangat terbantu dalam mengimplementasikan prinsip pembelajaran berdiferensiasi dan TaRL di kelas VII SMP."),
            ("28", "Merasa nyaman, tenang, dan percaya diri memanfaatkan teknologi AI sebagai sarana penunjang tugas profesional guru."),
            ("29", "Memiliki motivasi untuk terus menggunakan web application ini dalam pengembangan materi matematika pada topik lainnya."),
            ("30", "Bersedia merekomendasikan pemanfaatan web application ini kepada MGMP dan rekan sejawat guru matematika.")
        ])
    ]
    
    tbl_eval = doc.add_table(rows=1, cols=7)
    tbl_eval.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_eval, "94A3B8")
    col_w = [Inches(0.4), Inches(4.3), Inches(0.45), Inches(0.45), Inches(0.45), Inches(0.45), Inches(0.45)]
    headers = ["No.", "Butir Pernyataan Respon Guru", "STS", "TS", "CS", "S", "SS"]
    
    for c_idx, h in enumerate(headers):
        tbl_eval.rows[0].cells[c_idx].width = col_w[c_idx]
        format_cell(tbl_eval.rows[0].cells[c_idx], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color="F1F5F9")
        
    for group_title, items in items_data:
        hdr_row = tbl_eval.add_row()
        for w_idx, w in enumerate(col_w):
            hdr_row.cells[w_idx].width = w
        format_cell(hdr_row.cells[0], "", bold=True, size=10, color="E2E8F0")
        format_cell(hdr_row.cells[1], group_title, bold=True, size=10, color="E2E8F0")
        for c_idx in range(2, 7):
            format_cell(hdr_row.cells[c_idx], "", color="E2E8F0")
            
        for num, text in items:
            r = tbl_eval.add_row()
            for w_idx, w in enumerate(col_w):
                r.cells[w_idx].width = w
            format_cell(r.cells[0], num, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_cell(r.cells[1], text, size=10)
            for c_idx in range(2, 7):
                format_cell(r.cells[c_idx], "[  ]", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                
    # Bagian D: Kritik & Saran
    add_section_heading(doc, "D. Saran dan Masukan Kualitatif Guru")
    add_body_p(doc, "Tuliskan saran perbaikan, kendala yang dihadapi, atau harapan Bapak/Ibu untuk penyempurnaan web application ini ke depan:")
    tbl_notes = doc.add_table(rows=1, cols=1)
    tbl_notes.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_notes.rows[0].cells[0].width = Inches(7.1)
    format_cell(tbl_notes.rows[0].cells[0], "\n\n\n\n\n", size=11)
    set_table_borders(tbl_notes, "94A3B8")
    
    # Tanda tangan
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(8)
    
    tbl_sign = doc.add_table(rows=3, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sign, "FFFFFF")
    for r in tbl_sign.rows:
        r.cells[0].width = Inches(3.5)
        r.cells[1].width = Inches(3.5)
        
    format_cell(tbl_sign.rows[0].cells[0], "", size=11)
    format_cell(tbl_sign.rows[0].cells[1], "Tasikmalaya, .................................... 2026\nGuru Mata Pelajaran Matematika,", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_sign.rows[1].cells[0], "", size=11)
    format_cell(tbl_sign.rows[1].cells[1], "\n\n\n", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    format_cell(tbl_sign.rows[2].cells[0], "", size=11)
    format_cell(tbl_sign.rows[2].cells[1], "( .............................................................. )\nNIP/NUPTK. ...............................................", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    out_docx = os.path.join(OUT_DIR, "06_Angket_Kepraktisan_Guru.docx")
    doc.save(out_docx)
    print("Saved:", out_docx)

if __name__ == "__main__":
    build_instrument_06()

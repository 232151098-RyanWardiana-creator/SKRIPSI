import docx
import zipfile
import shutil
import os

DOCX_PATH = r"C:\Users\LATITUDE 3310 TOUCH\Documents\DOCUMENT RYAN\02_Kuliah\Semester_7\KF21518001 - Skripsi (Kelas A)\03_Draft Proposal (BAB 1-3)\Proposal_Skripsi_Ryan_Wardiana_Revisi_Bimbingan_2_Draft.docx"
IMAGE_PATH = r"C:\Users\LATITUDE 3310 TOUCH\Documents\DOCUMENT RYAN\02_Kuliah\Semester_7\KF21518001 - Skripsi (Kelas A)\03_Draft Proposal (BAB 1-3)\gambar_6_1_kerangka_teoretis.png"

def apply_revisions():
    doc = docx.Document(DOCX_PATH)
    body = doc._element.body

    # =========================================================================
    # 1. Definisi Operasional (DO) Heading in Body: GenAI -> Generative Artificial Intelligence (GenAI)
    # =========================================================================
    print("1. Updating Definisi Operasional 3.1 Heading...")
    # Cari heading 3.1
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Web Application LKPD Berbasis GenAI":
            print(f"   Found DO 3.1 heading at paragraph {i}")
            # Ganti teks heading menjadi lengkap dan tegak (regular)
            p.text = "Web Application LKPD Berbasis Generative Artificial Intelligence (GenAI)"
            p.style = doc.styles['Heading 2']
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = docx.shared.Pt(12)
                run.font.bold = True
                run.font.italic = False
            
            # Perbarui paragraf pertama di bawahnya jika mengandung teks awalan
            if i + 1 < len(doc.paragraphs):
                p_next = doc.paragraphs[i+1]
                if p_next.text.startswith("Web application LKPD berbasis GenAI"):
                    p_next.text = p_next.text.replace(
                        "Web application LKPD berbasis GenAI",
                        "Web application LKPD berbasis Generative Artificial Intelligence (GenAI)"
                    )
                    # Pastikan format font paragraf
                    for r in p_next.runs:
                        r.font.name = 'Times New Roman'
            break

    # =========================================================================
    # 2. Subbab 7.5: Hapus Validitas Kelayakan Produk Akbar 2013 & Rapikan
    # =========================================================================
    print("2. Restructuring Subbab 7.5 (Menghapus Kelayakan Produk Akbar)...")
    
    # Identifikasi elemen-elemen di Subbab 7.5
    idx_p168 = None
    for idx, el in enumerate(body):
        txt = ''.join(el.itertext()).strip()
        if txt.startswith("1. Analisis Kevalidan (Validitas Produk dan Instrumen)"):
            idx_p168 = idx
            break

    if idx_p168 is not None:
        print(f"   Found P168 at body index {idx_p168}")
        # body[idx_p168] -> Ubah judul menjadi: 1. Analisis Validitas Isi Instrumen Asesmen Diagnostik (Koefisien Aiken's V)
        p168_el = body[idx_p168]
        # Kosongkan runs dan isi baru
        for child in list(p168_el):
            if child.tag.endswith('r'):
                p168_el.remove(child)
        new_r = docx.oxml.OxmlElement('w:r')
        new_rPr = docx.oxml.OxmlElement('w:rPr')
        new_rFont = docx.oxml.OxmlElement('w:rFonts')
        new_rFont.set(docx.oxml.ns.qn('w:ascii'), 'Times New Roman')
        new_rFont.set(docx.oxml.ns.qn('w:hAnsi'), 'Times New Roman')
        new_b = docx.oxml.OxmlElement('w:b')
        new_sz = docx.oxml.OxmlElement('w:sz')
        new_sz.set(docx.oxml.ns.qn('w:val'), '24')
        new_rPr.append(new_rFont)
        new_rPr.append(new_b)
        new_rPr.append(new_sz)
        new_t = docx.oxml.OxmlElement('w:t')
        new_t.text = "1. Analisis Validitas Isi Instrumen Asesmen Diagnostik (Koefisien Aiken's V)"
        new_r.append(new_rPr)
        new_r.append(new_t)
        p168_el.append(new_r)

        # body[idx_p168 + 1] -> Narasi pengantar yang terfokus pada validitas isi instrumen
        p169_el = body[idx_p168 + 1]
        for child in list(p169_el):
            if child.tag.endswith('r'):
                p169_el.remove(child)
        r_p169 = docx.oxml.OxmlElement('w:r')
        rPr_169 = docx.oxml.OxmlElement('w:rPr')
        rFont_169 = docx.oxml.OxmlElement('w:rFonts')
        rFont_169.set(docx.oxml.ns.qn('w:ascii'), 'Times New Roman')
        rFont_169.set(docx.oxml.ns.qn('w:hAnsi'), 'Times New Roman')
        sz_169 = docx.oxml.OxmlElement('w:sz')
        sz_169.set(docx.oxml.ns.qn('w:val'), '24')
        rPr_169.append(rFont_169)
        rPr_169.append(sz_169)
        t_169 = docx.oxml.OxmlElement('w:t')
        t_169.text = (
            "Analisis kevalidan dilakukan untuk memastikan kelayakan teoretis dan empiris "
            "butir instrumen asesmen diagnostik materi rasio yang dikembangkan. Validitas isi "
            "setiap butir instrumen tes diagnostik dianalisis berdasarkan penilaian panel ahli "
            "validator menggunakan koefisien validitas isi Aiken's V (Aiken, 1985):"
        )
        r_p169.append(rPr_169)
        r_p169.append(t_169)
        p169_el.append(r_p169)

        # Elemen idx_p168 + 2 ("a. Validitas Isi...") dan idx_p168 + 3 ("Validitas isi setiap butir...") sekarang redundan
        # Kita hapus body[idx_p168 + 2] dan body[idx_p168 + 3]
        el_a = body[idx_p168 + 2]
        el_narr = body[idx_p168 + 3]
        body.remove(el_a)
        body.remove(el_narr)
        print("   Removed redundant sub-heading 'a.' and duplicate narrative")

    # Sekarang cari elemen "b. Validitas Kelayakan Produk LKPD..." s.d. Tabel 7.10 dan hapus seluruhnya!
    idx_b = None
    for idx, el in enumerate(body):
        txt = ''.join(el.itertext()).strip()
        if "b. Validitas Kelayakan Produk LKPD" in txt:
            idx_b = idx
            break

    if idx_b is not None:
        print(f"   Found 'b. Validitas Kelayakan Produk' at body index {idx_b}")
        # Hapus elemen mulai dari idx_b sampai tepat sebelum "2. Analisis Kepraktisan Penggunaan Produk"
        elements_to_remove = []
        curr_idx = idx_b
        while curr_idx < len(body):
            txt = ''.join(body[curr_idx].itertext()).strip()
            if txt.startswith("2. Analisis Kepraktisan Penggunaan Produk"):
                break
            elements_to_remove.append(body[curr_idx])
            curr_idx += 1

        for el in elements_to_remove:
            tag = el.tag.split('}')[-1]
            txt = ''.join(el.itertext())[:40].strip()
            print(f"      Removing [{tag}]: {txt}...")
            body.remove(el)
        print(f"   Successfully removed {len(elements_to_remove)} elements of Akbar product validity")

    # =========================================================================
    # 3. Penomoran Ulang Tabel: Tabel Kepraktisan -> Tabel 7.7, Matriks Jadwal -> Tabel 7.8
    # =========================================================================
    print("3. Renumbering Tables in Body (Tabel 7.7 & Tabel 7.8)...")
    for p in doc.paragraphs:
        # Perbarui rujukan di teks kepraktisan
        if "mengacu pada Tabel 7.11 berikut" in p.text:
            p.text = p.text.replace("mengacu pada Tabel 7.11 berikut", "mengacu pada Tabel 7.7 berikut")
            print("   Updated reference to Tabel 7.7 in practicality text")

        # Perbarui caption Tabel 7.11 -> Tabel 7.7
        if p.text.startswith("Tabel 7.11 Kriteria Interval Kepraktisan"):
            p.text = "Tabel 7.7 Kriteria Interval Kepraktisan dan Respons Pengguna (diadaptasi dari Riduwan, 2006, hlm. 89)"
            p.style = doc.styles['Caption']
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = docx.shared.Pt(10)
                r.font.bold = True
            print("   Renumbered caption to Tabel 7.7")

        # Perbarui rujukan di teks jadwal Subbab 7.6
        if "disajikan pada Tabel 7.12 berikut" in p.text:
            p.text = p.text.replace("disajikan pada Tabel 7.12 berikut", "disajikan pada Tabel 7.8 berikut")
            print("   Updated reference to Tabel 7.8 in schedule text")

        # Perbarui caption Tabel 7.12 -> Tabel 7.8
        if p.text.startswith("Tabel 7.12 Matriks Jadwal Rencana"):
            p.text = "Tabel 7.8 Matriks Jadwal Rencana Penelitian dan Pengembangan"
            p.style = doc.styles['Caption']
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = docx.shared.Pt(10)
                r.font.bold = True
            print("   Renumbered caption to Tabel 7.8")

    # =========================================================================
    # 4. Perbarui DAFTAR TABEL di Awal Dokumen
    # =========================================================================
    print("4. Updating DAFTAR TABEL...")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if "Tabel 7.10 Kriteria Interval Kevalidan Ahli" in txt:
            p.text = "Tabel 7.7 Kriteria Interval Kepraktisan dan Respons Pengguna\t33"
            print(f"   Updated DAFTAR TABEL item at P{i} to Tabel 7.7")
        elif "Tabel 7.11 Kriteria Interval Kepraktisan" in txt:
            p.text = "Tabel 7.8 Matriks Jadwal Rencana Penelitian dan Pengembangan\t35"
            print(f"   Updated DAFTAR TABEL item at P{i} to Tabel 7.8")
        elif "Tabel 7.12 Matriks Jadwal Rencana" in txt:
            # Hapus paragraf baris ke-37 yang kini redundan
            p_el = p._p
            p_el.getparent().remove(p_el)
            print(f"   Removed redundant DAFTAR TABEL item at P{i}")

    # Simpan dokumen
    doc.save(DOCX_PATH)
    print(f"Saved modified docx to {DOCX_PATH}")

    # =========================================================================
    # 5. Ganti Citra word/media/image2.png di dalam Arsip DOCX
    # =========================================================================
    print("5. Replacing image2.png with new Apple-style diagram in docx zip...")
    temp_zip = DOCX_PATH + ".temp.zip"
    shutil.copy2(DOCX_PATH, temp_zip)
    
    with zipfile.ZipFile(temp_zip, 'r') as zin:
        with zipfile.ZipFile(DOCX_PATH, 'w') as zout:
            for item in zin.infolist():
                if item.filename == "word/media/image2.png":
                    with open(IMAGE_PATH, 'rb') as f_img:
                        img_bytes = f_img.read()
                    zout.writestr(item, img_bytes)
                    print(f"   Replaced {item.filename} with {len(img_bytes):,} bytes from {os.path.basename(IMAGE_PATH)}")
                else:
                    zout.writestr(item, zin.read(item.filename))
    
    if os.path.exists(temp_zip):
        os.remove(temp_zip)
    print("   ZIP replacement complete.")

if __name__ == "__main__":
    apply_revisions()
